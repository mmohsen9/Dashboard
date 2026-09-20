# -*- coding: utf-8 -*-
"""
تطبيق الداشبورد التفاعلي - نظام إدارة المخزون والمبيعات (شركة الكوب الأول)
Interactive Inventory & Sales Dashboard — Streamlit + SQLite (RTL Arabic)

يحوّل السكربت الثابت إلى تطبيق تفاعلي:
- إدخال المبيعات اليومية وتحديث المخزون تلقائيًا
- تسجيل دفعات التحميص مع فحص نسبة الإنتاج (Yield) لحظيًا
- موازنة قابلة للتحرير مع نسبة نمو لكل حساب
- مركز تنبيهات + استيراد ملفاتك الحالية + تصدير Excel بنفس التنسيق

التشغيل:  streamlit run "تطبيق_الداشبورد_التفاعلي.py"
"""
from __future__ import annotations

import csv
import io
import re
import sqlite3
from contextlib import closing
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import streamlit as st

# ==========================================================================
# توافق مع إصدارات Streamlit القديمة (يعمل على 1.12 وأحدث)
# ==========================================================================
if not hasattr(st, "divider"):
    st.divider = lambda: st.markdown("---")

_st_dataframe = st.dataframe


def _df_to_html(df, hide_index):
    html = df.to_html(index=not hide_index, border=0, na_rep="", escape=True,
                      float_format=lambda x: f"{x:,.2f}", classes="rtl-tbl")
    return (
        "<style>.rtl-wrap{max-height:540px;overflow:auto;direction:rtl}"
        ".rtl-tbl{width:100%;border-collapse:collapse;font-size:0.9rem}"
        ".rtl-tbl th,.rtl-tbl td{border:1px solid rgba(128,128,128,.35);"
        "padding:6px 10px;text-align:right;white-space:nowrap}"
        ".rtl-tbl th{background:rgba(128,128,128,.18);position:sticky;top:0}</style>"
        "<div class='rtl-wrap'>" + html + "</div>"
    )


def _compat_df(target, orig, data=None, *args, **kwargs):
    """يحاول العرض الأصلي؛ وعند فشل الوسائط في الإصدارات القديمة يعرض جدول HTML
    داخل نفس الحاوية (target) — يعمل مع st و مع الأعمدة/التبويبات/الحاويات."""
    try:
        return orig(data, *args, **kwargs)
    except TypeError:
        if hasattr(data, "to_html"):
            try:
                return target.markdown(
                    _df_to_html(data, bool(kwargs.get("hide_index", False))),
                    unsafe_allow_html=True,
                )
            except Exception:
                pass
        for _k in ("use_container_width", "hide_index", "column_config", "width"):
            kwargs.pop(_k, None)
        return orig(data, *args, **kwargs)


st.dataframe = lambda data=None, *a, **k: _compat_df(st, _st_dataframe, data, *a, **k)

# ترقيع دالة الكلاس أيضًا حتى تشمل النداءات على الأعمدة/التبويبات مثل col.dataframe(...)
try:
    from streamlit.delta_generator import DeltaGenerator as _DG

    _dg_dataframe = _DG.dataframe

    def _dg_compat(self, data=None, *a, **k):
        return _compat_df(self, lambda d=None, *aa, **kk: _dg_dataframe(self, d, *aa, **kk),
                          data, *a, **k)

    _DG.dataframe = _dg_compat
except Exception:
    pass

# ==========================================================================
# المسارات والإعدادات الأساسية
# ==========================================================================
BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "01_ملفات_الإدخال_الشهرية"
STATIC_DIR = BASE_DIR / "02_مصادر_ثابتة"
DB_DIR = BASE_DIR / "03_قاعدة_البيانات"
OUTPUT_DIR = BASE_DIR / "04_المخرجات"

for _d in (INPUT_DIR, STATIC_DIR, DB_DIR, OUTPUT_DIR):
    _d.mkdir(parents=True, exist_ok=True)

DB_PATH = DB_DIR / "app_interactive.db"
DASHBOARD_PATH = OUTPUT_DIR / "الداشبورد_المالي_والتشغيلي.xlsx"
ALERTS_PATH = OUTPUT_DIR / "تنبيهات_الشهر.xlsx"

DEFAULT_SETTINGS = {
    "growth_rate": "0.10",      # نسبة النمو الافتراضية للموازنة
    "roast_yield": "0.833",     # نسبة إنتاج التحميص المتوقعة
    "roast_tolerance": "0.012", # التفاوت المسموح به
    "base_year": "2025",        # سنة الأساس (الفعلي)
    "budget_year": "2026",      # سنة الموازنة
}

SALE = "بيع"
GREEN_OUT = "خروج أخضر للتحميص"
ROASTED_IN = "دخول محمص من التحميص"
BRANCH_OUT = "صرف محمص للفروع"
SAMPLE = "عينات وهدايا"
ADJUST = "تسوية أو إعادة تقييم"
OTHER = "أخرى"

# ==========================================================================
# أدوات مساعدة (منقولة من السكربت الأصلي)
# ==========================================================================
def normalize(value):
    text = str(value).replace("\n", " ").replace("\r", " ").strip().lower()
    return re.sub(r"\s+", " ", text)


def clean_code(value):
    if pd.isna(value):
        return ""
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") else text


def number(value):
    if pd.isna(value):
        return 0.0
    text = str(value).strip().replace(",", "")
    text = text.replace("(", "-").replace(")", "")
    result = pd.to_numeric(text, errors="coerce")
    return float(result) if pd.notna(result) else 0.0


def find_column(columns, words):
    for column in columns:
        if any(word in normalize(column) for word in words):
            return column
    return None


def read_excel(path):
    try:
        return pd.read_excel(path, sheet_name=0)
    except Exception:
        return pd.DataFrame()


def read_csv_rows(path):
    for encoding in ["utf-8-sig", "utf-8", "cp1256", "latin1"]:
        try:
            with open(path, "r", encoding=encoding, errors="replace", newline="") as file:
                return list(csv.reader(file, delimiter=";"))
        except Exception:
            continue
    return []


def item_meta(code):
    """يحدد نوع البن ورقم العائلة من كود الصنف."""
    code = (code or "").upper().strip()
    if code.startswith("GB"):
        coffee_type = "بن أخضر"
    elif code.startswith("RB"):
        coffee_type = "بن محمص"
    else:
        coffee_type = "غير مصنف"
    family = re.sub(r"^(GB|RB)", "", code, flags=re.IGNORECASE)
    return coffee_type, family


# ==========================================================================
# طبقة قاعدة البيانات (SQLite)
# ==========================================================================
def connect():
    conn = sqlite3.connect(str(DB_PATH), check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn


def execute(sql, params=()):
    with closing(connect()) as conn:
        cur = conn.execute(sql, params)
        conn.commit()
        return cur.lastrowid


def executemany(sql, seq):
    with closing(connect()) as conn:
        conn.executemany(sql, seq)
        conn.commit()


def query_df(sql, params=()):
    with closing(connect()) as conn:
        return pd.read_sql_query(sql, conn, params=params)


def init_db():
    with closing(connect()) as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY, value TEXT
            );
            CREATE TABLE IF NOT EXISTS accounts (
                code TEXT PRIMARY KEY, name TEXT, type TEXT
            );
            CREATE TABLE IF NOT EXISTS items (
                code TEXT PRIMARY KEY, name TEXT, coffee_type TEXT, family TEXT
            );
            CREATE TABLE IF NOT EXISTS movements (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                year INTEGER, month INTEGER, mdate TEXT,
                item_code TEXT, item_name TEXT, coffee_type TEXT, family TEXT,
                qty REAL, value REAL, total REAL,
                movement_type TEXT, source TEXT, reference TEXT,
                contact TEXT, description TEXT, branch TEXT,
                revenue_account TEXT, inventory_account TEXT,
                created_at TEXT, created_by TEXT, from_report INTEGER DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS item_balances (
                item_code TEXT PRIMARY KEY, item_name TEXT, coffee_type TEXT, family TEXT,
                qty REAL, value REAL, year INTEGER, month INTEGER,
                source_file TEXT, updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS branch_sales (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                year INTEGER, month INTEGER, mdate TEXT,
                invoice TEXT, contact TEXT, reference TEXT,
                amount REAL, status TEXT, source TEXT, branch TEXT,
                source_file TEXT, imported_at TEXT
            );
            CREATE TABLE IF NOT EXISTS channel_sales (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                year INTEGER, month INTEGER, channel TEXT,
                amount REAL, note TEXT, created_at TEXT, created_by TEXT,
                UNIQUE(year, month, channel)
            );
            CREATE TABLE IF NOT EXISTS pnl_monthly (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                year INTEGER, month INTEGER,
                account_code TEXT, account_name TEXT, category TEXT,
                amount REAL, source_file TEXT, imported_at TEXT
            );
            CREATE TABLE IF NOT EXISTS fin_figures (
                year INTEGER PRIMARY KEY,
                current_assets REAL, inventory REAL, current_liabilities REAL,
                total_assets REAL, total_equity REAL,
                source_file TEXT, updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS financial_lines (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                statement TEXT, year INTEGER, month INTEGER,
                account_code TEXT, account_name TEXT, account_type TEXT,
                amount REAL
            );
            CREATE TABLE IF NOT EXISTS raw_statements (
                statement TEXT, year INTEGER, month INTEGER,
                payload TEXT, imported_at TEXT
            );
            CREATE TABLE IF NOT EXISTS budget (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                year INTEGER, month INTEGER,
                account_code TEXT, account_name TEXT, account_type TEXT,
                actual_prev REAL, growth REAL, proposed REAL, approved REAL
            );
            CREATE TABLE IF NOT EXISTS alerts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT, severity TEXT, category TEXT,
                message TEXT, resolved INTEGER DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS audit_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                ts TEXT, usr TEXT, action TEXT, details TEXT
            );
            """
        )
        # ترحيل: إضافة العمود لقواعد البيانات القديمة (نقطة جديدة، غير مخلّة بالبنية)
        try:
            conn.execute("ALTER TABLE movements ADD COLUMN from_report INTEGER DEFAULT 0")
        except Exception:
            pass
        # ترحيل لمرة واحدة لميزة الأرصدة المعتمدة: تنظيف الحركات المستوردة القديمة
        # (كانت غير موسومة فتُحسب خطأً كإدخال يدوي)؛ يُعاد ملؤها عند أول استيراد.
        row = conn.execute("SELECT value FROM settings WHERE key='schema_version'").fetchone()
        version = row[0] if row else "1"
        if str(version) < "2":
            conn.execute("DELETE FROM movements")
            conn.execute("DELETE FROM item_balances")
            conn.execute("DELETE FROM branch_sales")
            conn.execute(
                "INSERT INTO settings(key, value) VALUES('schema_version', '2') "
                "ON CONFLICT(key) DO UPDATE SET value='2'"
            )
        for key, val in DEFAULT_SETTINGS.items():
            conn.execute("INSERT OR IGNORE INTO settings(key, value) VALUES(?, ?)", (key, val))
        conn.commit()


# ---- الإعدادات ----
def get_settings():
    df = query_df("SELECT key, value FROM settings")
    data = dict(zip(df["key"], df["value"])) if not df.empty else {}
    merged = dict(DEFAULT_SETTINGS)
    merged.update(data)
    return merged


def set_setting(key, value):
    execute(
        "INSERT INTO settings(key, value) VALUES(?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, str(value)),
    )


def s_float(settings, key):
    try:
        return float(settings.get(key, DEFAULT_SETTINGS[key]))
    except Exception:
        return float(DEFAULT_SETTINGS[key])


def s_int(settings, key):
    try:
        return int(float(settings.get(key, DEFAULT_SETTINGS[key])))
    except Exception:
        return int(float(DEFAULT_SETTINGS[key]))


# ---- سجل التدقيق ----
def audit(action, details=""):
    execute(
        "INSERT INTO audit_log(ts, usr, action, details) VALUES(?, ?, ?, ?)",
        (now_str(), current_user(), action, details),
    )


def now_str():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def current_user():
    try:
        return st.session_state.get("user_name", "admin")
    except Exception:
        return "admin"


# ---- الأصناف والحسابات ----
def upsert_item(code, name, coffee_type=None, family=None):
    code = (code or "").upper().strip()
    if not code:
        return
    if coffee_type is None or family is None:
        coffee_type, family = item_meta(code)
    execute(
        "INSERT INTO items(code, name, coffee_type, family) VALUES(?, ?, ?, ?) "
        "ON CONFLICT(code) DO UPDATE SET name=excluded.name, "
        "coffee_type=excluded.coffee_type, family=excluded.family",
        (code, str(name or "").strip(), coffee_type, family),
    )


def upsert_account(code, name, acc_type=""):
    code = clean_code(code)
    if not code:
        return
    execute(
        "INSERT INTO accounts(code, name, type) VALUES(?, ?, ?) "
        "ON CONFLICT(code) DO UPDATE SET name=excluded.name, type=excluded.type",
        (code, str(name or "").strip(), str(acc_type or "").strip()),
    )


def list_items():
    return query_df("SELECT code, name, coffee_type, family FROM items ORDER BY code")


def list_accounts():
    return query_df("SELECT code, name, type FROM accounts ORDER BY code")


# ---- الحركات ----
MOVE_COLS = [
    "year", "month", "mdate", "item_code", "item_name", "coffee_type", "family",
    "qty", "value", "total", "movement_type", "source", "reference",
    "contact", "description", "branch", "revenue_account", "inventory_account",
    "created_at", "created_by",
]


def add_movement(**kw):
    code = (kw.get("item_code") or "").upper().strip()
    coffee_type, family = item_meta(code)
    row = {c: kw.get(c) for c in MOVE_COLS}
    row["item_code"] = code
    row["coffee_type"] = kw.get("coffee_type") or coffee_type
    row["family"] = kw.get("family") or family
    row["created_at"] = now_str()
    row["created_by"] = current_user()
    row["qty"] = float(kw.get("qty") or 0)
    row["value"] = float(kw.get("value") or 0)
    row["total"] = float(kw.get("total") or 0)
    execute(
        "INSERT INTO movements(" + ",".join(MOVE_COLS) + ") VALUES(" + ",".join(["?"] * len(MOVE_COLS)) + ")",
        tuple(row.get(c) for c in MOVE_COLS),
    )
    if row["item_name"]:
        upsert_item(code, row["item_name"], row["coffee_type"], row["family"])


def current_stock():
    """الرصيد = الرصيد الختامي المعتمد من تقرير حركة المخزون + الحركات اليدوية اللاحقة فقط.
    (لا نجمع حركات التقرير لأنها لا تطابق الرصيد الختامي؛ التقرير هو المرجع)."""
    bal = query_df(
        "SELECT item_code, item_name, coffee_type, qty, value FROM item_balances"
    )
    man = query_df(
        "SELECT item_code, MAX(item_name) AS item_name, MAX(coffee_type) AS coffee_type, "
        "COALESCE(SUM(qty),0) AS qty, COALESCE(SUM(value),0) AS value "
        "FROM movements WHERE COALESCE(from_report,0)=0 AND item_code<>'' GROUP BY item_code"
    )
    cols = ["كود الصنف", "اسم الصنف", "نوع البن", "الرصيد الحالي", "قيمة الرصيد"]
    if bal.empty and man.empty:
        return pd.DataFrame(columns=cols)
    merged = pd.merge(bal, man, on="item_code", how="outer", suffixes=("_b", "_m"))
    for c in ["qty_b", "qty_m", "value_b", "value_m"]:
        if c not in merged:
            merged[c] = 0
    merged = merged.fillna({"qty_b": 0, "qty_m": 0, "value_b": 0, "value_m": 0})
    out = pd.DataFrame({
        "كود الصنف": merged["item_code"],
        "اسم الصنف": merged.get("item_name_b").fillna(merged.get("item_name_m")).fillna(""),
        "نوع البن": merged.get("coffee_type_b").fillna(merged.get("coffee_type_m")).fillna(""),
        "الرصيد الحالي": (merged["qty_b"] + merged["qty_m"]).round(2),
        "قيمة الرصيد": (merged["value_b"] + merged["value_m"]).round(2),
    })
    return out[out["كود الصنف"].astype(str) != ""].sort_values("كود الصنف").reset_index(drop=True)


def stock_of(item_code):
    b = query_df("SELECT COALESCE(qty,0) AS s FROM item_balances WHERE item_code=?", (item_code,))
    base = float(b["s"].iloc[0]) if not b.empty else 0.0
    m = query_df(
        "SELECT COALESCE(SUM(qty),0) AS s FROM movements "
        "WHERE item_code=? AND COALESCE(from_report,0)=0", (item_code,)
    )
    manual = float(m["s"].iloc[0]) if not m.empty else 0.0
    return base + manual


def movements_df():
    """يرجع الحركات بأعمدة عربية تتوافق مع دوال الحساب."""
    df = query_df(
        "SELECT year AS 'السنة', month AS 'الشهر', mdate AS 'التاريخ', "
        "item_code AS 'كود الصنف', item_name AS 'اسم الصنف', coffee_type AS 'نوع البن', "
        "family AS 'رقم العائلة', qty AS 'حركة الكمية', value AS 'قيمة الحركة', "
        "total AS 'الإجمالي', movement_type AS 'نوع الحركة', source AS 'المصدر', "
        "reference AS 'المرجع', contact AS 'العميل أو المورد', description AS 'الوصف', "
        "branch AS 'الفرع' FROM movements ORDER BY id"
    )
    return df


def movement_years():
    df = query_df("SELECT DISTINCT year FROM movements WHERE year IS NOT NULL ORDER BY year")
    return [int(y) for y in df["year"].tolist()] if not df.empty else []


def movement_months(year):
    df = query_df(
        "SELECT DISTINCT month FROM movements WHERE year=? AND month IS NOT NULL ORDER BY month",
        (int(year),),
    )
    return [int(m) for m in df["month"].tolist()] if not df.empty else []


def movement_period_summary(year=None, month=None):
    """ملخّص حركة المخزون لكل صنف خلال فترة (سنة/شهر) بناءً على مدخلات الإكسيل:
    الوارد (كميات موجبة) / المنصرف (كميات سالبة) / صافي الحركة / قيمة الحركة."""
    where = ["item_code<>''"]
    params = []
    if year is not None:
        where.append("year=?")
        params.append(int(year))
    if month is not None:
        where.append("month=?")
        params.append(int(month))
    sql = (
        "SELECT item_code, MAX(item_name) AS item_name, MAX(coffee_type) AS coffee_type, "
        "COALESCE(SUM(CASE WHEN qty>0 THEN qty ELSE 0 END),0) AS incoming, "
        "COALESCE(SUM(CASE WHEN qty<0 THEN -qty ELSE 0 END),0) AS outgoing, "
        "COALESCE(SUM(qty),0) AS net, COALESCE(SUM(value),0) AS val "
        "FROM movements WHERE " + " AND ".join(where) + " GROUP BY item_code ORDER BY item_code"
    )
    df = query_df(sql, tuple(params))
    if df.empty:
        return df
    df = df.rename(columns={
        "item_code": "كود الصنف", "item_name": "اسم الصنف", "coffee_type": "نوع البن",
        "incoming": "الوارد", "outgoing": "المنصرف", "net": "صافي الحركة", "val": "قيمة الحركة",
    })
    return df.round({"الوارد": 2, "المنصرف": 2, "صافي الحركة": 2, "قيمة الحركة": 2})


# فئات الحركة المعروضة في البيان المجمّع للصنف (ترتيب منطقي: وارد ← تحميص ← منصرف ← تسوية)
STATEMENT_CATEGORIES = [
    ("شراء / وارد", OTHER),
    ("تحميص (دخول محمص)", ROASTED_IN),
    ("تحميص (خروج أخضر)", GREEN_OUT),
    ("بيع", SALE),
    ("صرف للفروع", BRANCH_OUT),
    ("عينات وهدايا", SAMPLE),
    ("فروقات جردية", ADJUST),
]


def item_movement_statement(year, month=None):
    """بيان مجمّع لكل صنف من سجل الحركات: مجموع الكمية لكل فئة (تحميص/بيع/شراء/فروقات جردية/
    صرف للفروع/عينات) + صافي الحركة، مفلتراً على مستوى العام والشهر."""
    where = ["item_code<>''", "year=?"]
    params = [int(year)]
    if month is not None:
        where.append("month=?")
        params.append(int(month))
    sel = ["item_code", "MAX(item_name) AS item_name"]
    cat_params = []
    for i, (_label, mt) in enumerate(STATEMENT_CATEGORIES):
        sel.append(f"COALESCE(SUM(CASE WHEN movement_type=? THEN qty ELSE 0 END),0) AS c{i}")
        cat_params.append(mt)
    sel.append("COALESCE(SUM(qty),0) AS net")
    sql = ("SELECT " + ", ".join(sel) + " FROM movements WHERE " + " AND ".join(where) +
           " GROUP BY item_code ORDER BY item_code")
    # ترتيب الوسائط: وسائط CASE في SELECT أولاً ثم وسائط WHERE
    df = query_df(sql, tuple(cat_params) + tuple(params))
    if df.empty:
        return df
    rename = {"item_code": "كود الصنف", "item_name": "اسم الصنف", "net": "صافي الحركة"}
    for i, (label, _mt) in enumerate(STATEMENT_CATEGORIES):
        rename[f"c{i}"] = label
    df = df.rename(columns=rename)
    cat_labels = [c[0] for c in STATEMENT_CATEGORIES]
    df = df[["كود الصنف", "اسم الصنف"] + cat_labels + ["صافي الحركة"]]
    df = df[df[cat_labels].abs().sum(axis=1) > 0].copy()  # أصناف لها حركة فقط
    df["_a"] = df[cat_labels].abs().sum(axis=1)
    df = df.sort_values("_a", ascending=False).drop(columns="_a").reset_index(drop=True)
    return df.round({c: 1 for c in cat_labels + ["صافي الحركة"]})


# ---- التنبيهات ----
def add_alert(severity, category, message):
    execute(
        "INSERT INTO alerts(created_at, severity, category, message, resolved) VALUES(?, ?, ?, ?, 0)",
        (now_str(), severity, category, message),
    )


def open_alerts_count():
    df = query_df("SELECT COUNT(*) AS c FROM alerts WHERE resolved=0")
    return int(df["c"].iloc[0]) if not df.empty else 0


# ==========================================================================
# دوال الحساب (منقولة ومكيّفة لتقرأ من قاعدة البيانات)
# ==========================================================================
def pnl_df():
    return query_df(
        "SELECT year AS 'السنة', month AS 'الشهر', account_code AS 'كود الحساب', "
        "account_name AS 'اسم الحساب', account_type AS 'نوع الحساب', amount AS 'المبلغ' "
        "FROM financial_lines WHERE statement='pnl'"
    )


def build_budget(pnl, base_year, growth, budget_year):
    if pnl.empty:
        return pd.DataFrame()
    result = pnl[pnl["السنة"] == base_year].groupby(
        ["الشهر", "كود الحساب", "اسم الحساب", "نوع الحساب"], as_index=False
    )["المبلغ"].sum()
    if result.empty:
        return result
    result["السنة"] = budget_year
    result["actual_prev"] = result["المبلغ"]
    result["growth"] = growth
    result["proposed"] = result["actual_prev"] * (1 + growth)
    result["approved"] = result["proposed"]
    return result


def build_forecast(pnl):
    if pnl.empty:
        return pd.DataFrame()
    rows = []
    for (code, name, atype), group in pnl.groupby(["كود الحساب", "اسم الحساب", "نوع الحساب"]):
        values = group.sort_values(["السنة", "الشهر"])["المبلغ"]
        rows.append({
            "كود الحساب": code,
            "اسم الحساب": name,
            "نوع الحساب": atype,
            "متوسط آخر 6 أشهر": round(values.tail(6).mean(), 2),
            "متوسط آخر 12 شهرًا": round(values.tail(12).mean(), 2),
            "توقع الشهر التالي": round(values.tail(6).mean(), 2),
        })
    return pd.DataFrame(rows)


def build_budget_from_monthly(base_year, growth, budget_year):
    """موازنة شهرية مبنية على الفعلي الشهري (pnl_monthly) لسنة الأساس:
    لكل حساب ولكل شهر: مقترح = فعلي نفس الشهر × (1 + النمو)."""
    df = query_df(
        "SELECT month, account_code, account_name, category, COALESCE(SUM(amount),0) AS amt "
        "FROM pnl_monthly WHERE year=? GROUP BY month, account_code, account_name, category",
        (int(base_year),),
    )
    if df.empty:
        return pd.DataFrame()
    df = df.rename(columns={"month": "الشهر", "account_code": "كود الحساب",
                            "account_name": "اسم الحساب", "category": "نوع الحساب", "amt": "المبلغ"})
    df["السنة"] = int(budget_year)
    df["actual_prev"] = df["المبلغ"].astype(float)
    df["growth"] = float(growth)
    df["proposed"] = df["actual_prev"] * (1 + float(growth))
    df["approved"] = df["proposed"]
    return df


def build_forecast_monthly():
    """توقع لكل حساب = متوسط آخر 6/12 شهرًا من الفعلي الشهري عبر كل الأعوام."""
    df = query_df(
        "SELECT year, month, account_code, account_name, COALESCE(SUM(amount),0) AS amt "
        "FROM pnl_monthly GROUP BY year, month, account_code, account_name "
        "ORDER BY account_code, year, month"
    )
    if df.empty:
        return pd.DataFrame()
    rows = []
    for (code, name), g in df.groupby(["account_code", "account_name"]):
        vals = g.sort_values(["year", "month"])["amt"]
        rows.append({
            "كود الحساب": code, "اسم الحساب": name,
            "متوسط آخر 6 أشهر": round(float(vals.tail(6).mean()), 2),
            "متوسط آخر 12 شهرًا": round(float(vals.tail(12).mean()), 2),
            "توقع الشهر التالي": round(float(vals.tail(6).mean()), 2),
        })
    return pd.DataFrame(rows)


def build_roasting(inventory, roast_yield, tolerance):
    if inventory.empty:
        return pd.DataFrame()
    green = inventory[inventory["نوع الحركة"] == GREEN_OUT]
    roasted = inventory[inventory["نوع الحركة"] == ROASTED_IN]
    if green.empty and roasted.empty:
        return pd.DataFrame()

    green_summary = green.groupby(["السنة", "الشهر", "رقم العائلة"], as_index=False)["حركة الكمية"].sum()
    green_summary = green_summary.rename(columns={"حركة الكمية": "كمية الأخضر المحولة"})
    green_summary["كمية الأخضر المحولة"] = green_summary["كمية الأخضر المحولة"].abs()

    roasted_summary = roasted.groupby(["السنة", "الشهر", "رقم العائلة"], as_index=False)["حركة الكمية"].sum()
    roasted_summary = roasted_summary.rename(columns={"حركة الكمية": "كمية المحمص الناتجة"})

    result = green_summary.merge(roasted_summary, on=["السنة", "الشهر", "رقم العائلة"], how="outer").fillna(0)
    if result.empty:
        return result

    names = inventory[["رقم العائلة", "اسم الصنف"]].copy()
    names["اسم الصنف"] = names["اسم الصنف"].astype(str).str.strip()
    names = names[names["اسم الصنف"] != ""].drop_duplicates("رقم العائلة").rename(columns={"اسم الصنف": "الصنف"})
    result = result.merge(names, on="رقم العائلة", how="left")

    result["الناتج المتوقع"] = result["كمية الأخضر المحولة"] * roast_yield
    result["فرق الكمية"] = result["كمية المحمص الناتجة"] - result["الناتج المتوقع"]
    result["فرق النسبة"] = result.apply(
        lambda r: r["فرق الكمية"] / r["الناتج المتوقع"] if r["الناتج المتوقع"] else 0, axis=1
    )
    result["الحالة"] = result["فرق النسبة"].map(
        lambda v: "مقبول" if abs(v) <= tolerance else "يحتاج مراجعة"
    )
    return result[[
        "السنة", "الشهر", "رقم العائلة", "الصنف", "كمية الأخضر المحولة",
        "كمية المحمص الناتجة", "الناتج المتوقع", "فرق الكمية", "فرق النسبة", "الحالة",
    ]]


def build_inventory_summary(inventory):
    if inventory.empty:
        return pd.DataFrame()
    return inventory.groupby(
        ["السنة", "الشهر", "كود الصنف", "اسم الصنف", "نوع البن", "رقم العائلة", "نوع الحركة"],
        as_index=False,
    ).agg({"حركة الكمية": "sum", "قيمة الحركة": "sum", "الإجمالي": "sum"})


def branch_sales_df():
    return query_df(
        "SELECT year AS 'السنة', month AS 'الشهر', mdate AS 'التاريخ', invoice AS 'رقم الفاتورة', "
        "contact AS 'العميل', reference AS 'المرجع', amount AS 'المبلغ', status AS 'الحالة', "
        "branch AS 'الفرع' FROM branch_sales ORDER BY mdate, invoice"
    )


def branch_sales_summary():
    df = branch_sales_df()
    if df.empty:
        return df, pd.DataFrame()
    by_branch = df.groupby("الفرع", as_index=False).agg(
        **{"عدد الفواتير": ("رقم الفاتورة", "count"), "إجمالي المبيعات": ("المبلغ", "sum")}
    ).sort_values("إجمالي المبيعات", ascending=False).reset_index(drop=True)
    return df, by_branch


def branch_sales_grouped(year, month=None):
    """مبيعات مجمّعة من تقارير الإكسيل حسب العميل (Contact) والفرع (Branches)،
    على أساس سنوي (كل الأشهر) أو شهري (شهر محدّد)."""
    params = [int(year)]
    where = "year=?"
    if month:
        where += " AND month=?"
        params.append(int(month))
    df = query_df(
        "SELECT contact AS 'العميل', branch AS 'الفرع', COUNT(*) AS 'عدد الفواتير', "
        f"COALESCE(SUM(amount),0) AS 'إجمالي المبيعات' FROM branch_sales WHERE {where} "
        "GROUP BY contact, branch ORDER BY 4 DESC",
        tuple(params),
    )
    if not df.empty:
        df["العميل"] = df["العميل"].apply(
            lambda x: str(x).strip() if str(x).strip() and str(x).strip().lower() != "nan" else "(بدون عميل)")
        df["الفرع"] = df["الفرع"].apply(
            lambda x: str(x).strip() if str(x).strip() and str(x).strip().lower() != "nan" else "غير محدد")
    return df


# ---- مبيعات القنوات اليدوية (تطبيقات التوصيل + الموقع) ----
DELIVERY_CHANNELS = ["جاهز", "نوفمبر", "هنجرستيشن", "الموقع الإلكتروني", "أخرى"]


def upsert_channel_sale(year, month, channel, amount, note=""):
    execute(
        "INSERT INTO channel_sales(year, month, channel, amount, note, created_at, created_by) "
        "VALUES(?,?,?,?,?,?,?) ON CONFLICT(year, month, channel) DO UPDATE SET "
        "amount=excluded.amount, note=excluded.note, created_at=excluded.created_at",
        (int(year), int(month), channel, float(amount), note, now_str(), current_user()),
    )


def channel_sales_df():
    return query_df(
        "SELECT year AS 'السنة', month AS 'الشهر', channel AS 'القناة', "
        "amount AS 'المبلغ', note AS 'ملاحظة' FROM channel_sales ORDER BY year, month, channel"
    )


def list_years():
    years = set()
    for sql in (
        "SELECT DISTINCT year FROM pnl_monthly",
        "SELECT DISTINCT year FROM branch_sales",
        "SELECT DISTINCT year FROM channel_sales",
        "SELECT DISTINCT year FROM fin_figures",
    ):
        df = query_df(sql)
        if not df.empty:
            years.update(int(y) for y in df["year"].dropna().tolist())
    return sorted(years)


def monthly_sales_series(year):
    """مبيعات كل شهر (1..12) = فواتير الفروع + مبيعات القنوات اليدوية."""
    s = pd.Series(0.0, index=range(1, 13))
    inv = query_df(
        "SELECT month, COALESCE(SUM(amount),0) AS v FROM branch_sales WHERE year=? GROUP BY month",
        (int(year),),
    )
    for _, r in inv.iterrows():
        m = int(r["month"]) if pd.notna(r["month"]) else 0
        if 1 <= m <= 12:
            s[m] += float(r["v"])
    ch = query_df(
        "SELECT month, COALESCE(SUM(amount),0) AS v FROM channel_sales WHERE year=? GROUP BY month",
        (int(year),),
    )
    for _, r in ch.iterrows():
        m = int(r["month"]) if pd.notna(r["month"]) else 0
        if 1 <= m <= 12:
            s[m] += float(r["v"])
    return s


def monthly_pnl_series(year, category):
    s = pd.Series(0.0, index=range(1, 13))
    df = query_df(
        "SELECT month, COALESCE(SUM(amount),0) AS v FROM pnl_monthly "
        "WHERE year=? AND category=? GROUP BY month", (int(year), category),
    )
    for _, r in df.iterrows():
        m = int(r["month"]) if pd.notna(r["month"]) else 0
        if 1 <= m <= 12:
            s[m] = float(r["v"])
    return s


def pnl_year_summary(year):
    df = query_df(
        "SELECT category, COALESCE(SUM(amount),0) AS v FROM pnl_monthly WHERE year=? GROUP BY category",
        (int(year),),
    )
    d = {row["category"]: float(row["v"]) for _, row in df.iterrows()}
    revenue = d.get("revenue", 0.0)
    cogs = d.get("cogs", 0.0)
    expense = d.get("expense", 0.0)
    other = d.get("other_income", 0.0)
    gross = revenue - cogs
    net = gross - expense + other
    return {
        "الإيرادات": revenue, "تكلفة المبيعات": cogs, "مجمل الربح": gross,
        "المصاريف التشغيلية": expense, "إيرادات أخرى": other, "صافي الربح": net,
        "هامش مجمل الربح %": (gross / revenue * 100) if revenue else 0.0,
        "هامش صافي الربح %": (net / revenue * 100) if revenue else 0.0,
    }


def stock_value_total():
    df = query_df("SELECT COALESCE(SUM(value),0) AS v FROM item_balances")
    return float(df["v"].iloc[0]) if not df.empty else 0.0


def financial_ratios(year):
    """نِسب مالية لسنة: السيولة، السريعة، دوران المخزون، الهوامش، العائد."""
    fig = query_df(
        "SELECT current_assets, inventory, current_liabilities, total_assets, total_equity "
        "FROM fin_figures WHERE year=?", (int(year),),
    )
    ca = cl = inv = ta = te = 0.0
    if not fig.empty:
        ca = float(fig["current_assets"].iloc[0] or 0)
        inv = float(fig["inventory"].iloc[0] or 0)
        cl = float(fig["current_liabilities"].iloc[0] or 0)
        ta = float(fig["total_assets"].iloc[0] or 0)
        te = float(fig["total_equity"].iloc[0] or 0)
    if inv <= 0:
        inv = stock_value_total()
    pnl = pnl_year_summary(year)
    cogs, net, rev = pnl["تكلفة المبيعات"], pnl["صافي الربح"], pnl["الإيرادات"]
    turnover = (cogs / inv) if inv else 0.0
    ratios = {
        "النسبة الجارية (السيولة)": (ca / cl) if cl else 0.0,
        "نسبة السيولة السريعة": ((ca - inv) / cl) if cl else 0.0,
        "معدل دوران المخزون (مرة)": turnover,
        "متوسط فترة بقاء المخزون (يوم)": (365 / turnover) if turnover else 0.0,
        "هامش مجمل الربح %": pnl["هامش مجمل الربح %"],
        "هامش صافي الربح %": pnl["هامش صافي الربح %"],
        "العائد على الأصول % (ROA)": (net / ta * 100) if ta else 0.0,
        "العائد على حقوق الملكية % (ROE)": (net / te * 100) if te else 0.0,
    }
    figures = {"الأصول المتداولة": ca, "المخزون": inv, "الخصوم المتداولة": cl,
               "إجمالي الأصول": ta, "حقوق الملكية": te}
    return ratios, figures


# ==========================================================================
# الاستيراد من ملفاتك الحالية (منقول من السكربت الأصلي)
# ==========================================================================
def _read_chart_table(path):
    """يكتشف صف العناوين ديناميكيًا ويطابق الأعمدة بالاسم الدقيق
    (يعالج خطأ دليل الحسابات: العناوين ليست في الصف الأول + تطابق Account مع Account Code)."""
    try:
        raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
    except Exception:
        return None
    if raw.empty:
        return None
    header_idx = None
    for i in range(len(raw)):
        cells = [normalize(c) for c in raw.iloc[i].tolist()]
        if "account code" in cells or "كود الحساب" in cells or "رقم الحساب" in cells:
            header_idx = i
            break
    if header_idx is None:
        return None
    headers = [str(c).strip() for c in raw.iloc[header_idx].tolist()]
    body = raw.iloc[header_idx + 1:].copy()
    body.columns = headers

    def pick(exact):
        for col in headers:
            if normalize(col) in exact:
                return col
        return None

    code_col = pick(["account code", "رقم الحساب", "كود الحساب"])
    name_col = pick(["account", "account name", "اسم الحساب"])
    type_col = pick(["account type", "نوع الحساب"])
    if code_col is None or name_col is None:
        return None
    chart = pd.DataFrame({
        "كود الحساب": body[code_col].map(clean_code),
        "اسم الحساب": body[name_col].astype(str).str.strip(),
        "نوع الحساب": body[type_col].astype(str).str.strip() if type_col else "",
    })
    chart = chart[(chart["كود الحساب"] != "") & (chart["كود الحساب"].str.lower() != "nan")]
    return chart.drop_duplicates("كود الحساب")


def load_chart_of_accounts():
    files = list(STATIC_DIR.glob("*.xlsx")) + list(BASE_DIR.glob("*.xlsx"))
    for path in files:
        name = path.name.lower()
        if "chart" not in name and "account" not in name:
            continue
        chart = _read_chart_table(path)
        if chart is not None and not chart.empty:
            return chart, path.name
    return pd.DataFrame(columns=["كود الحساب", "اسم الحساب", "نوع الحساب"]), None


def attach_accounts(data, chart, code_column=None, name_column=None):
    output = data.copy()
    chart_by_code, chart_by_name = {}, {}
    for _, row in chart.iterrows():
        rec = {
            "كود الحساب": clean_code(row["كود الحساب"]),
            "اسم الحساب": str(row["اسم الحساب"]).strip(),
            "نوع الحساب": str(row["نوع الحساب"]).strip(),
        }
        chart_by_code[rec["كود الحساب"]] = rec
        chart_by_name[normalize(rec["اسم الحساب"])] = rec

    if code_column and code_column in output.columns:
        output["كود الحساب"] = output[code_column].map(clean_code)
    else:
        output["كود الحساب"] = ""
    if name_column and name_column in output.columns:
        source_names = output[name_column].astype(str).str.strip()
    else:
        source_names = pd.Series([""] * len(output), index=output.index)

    output["اسم الحساب"] = ""
    output["نوع الحساب"] = ""
    for index in output.index:
        code = clean_code(output.at[index, "كود الحساب"])
        source_name = str(source_names.loc[index]).strip()
        matched = chart_by_code.get(code)
        if matched is None and source_name:
            matched = chart_by_name.get(normalize(source_name))
        if matched is None and source_name:
            key = normalize(source_name)
            for cname, crec in chart_by_name.items():
                if key in cname or cname in key:
                    matched = crec
                    break
        if matched is not None:
            output.at[index, "كود الحساب"] = matched["كود الحساب"]
            output.at[index, "اسم الحساب"] = matched["اسم الحساب"]
            output.at[index, "نوع الحساب"] = matched["نوع الحساب"]
        else:
            output.at[index, "اسم الحساب"] = source_name
    return output


def extract_year_month(folder):
    match = re.search(r"(20\d{2})[-_](0?[1-9]|1[0-2])", str(folder))
    return (int(match.group(1)), int(match.group(2))) if match else (2025, 12)


def parse_profit_and_loss(data, year, chart):
    if data.empty:
        return pd.DataFrame()
    account_column = find_column(data.columns, ["account", "account name", "اسم الحساب"]) or data.columns[0]
    rows = []
    for _, row in data.iterrows():
        account_text = str(row.get(account_column, "")).strip()
        if not account_text:
            continue
        if normalize(account_text).startswith((
            "total", "gross profit", "net profit", "trading income", "cost of sales",
            "operating expenses", "other income", "إجمالي", "صافي",
        )):
            continue
        match = re.match(r"^\s*(\d+)\s*[-–]\s*(.*)$", account_text)
        source_code = match.group(1).strip() if match else ""
        source_name = match.group(2).strip() if match else account_text
        amount = sum(number(row.get(c, 0)) for c in data.columns if c != account_column)
        if amount:
            rows.append({
                "السنة": year, "الشهر": 12, "كود_مصدر": source_code,
                "اسم_مصدر": source_name, "المبلغ": amount,
            })
    result = pd.DataFrame(rows)
    if result.empty:
        return result
    result = attach_accounts(result, chart, "كود_مصدر", "اسم_مصدر")
    return result[["السنة", "الشهر", "كود الحساب", "اسم الحساب", "نوع الحساب", "المبلغ"]]


def parse_inventory(path, year, month):
    rows = read_csv_rows(path)
    if not rows:
        return pd.DataFrame()
    header_index = None
    for index, row in enumerate(rows):
        text = normalize(" ".join(row))
        if "date" in text and "qoh movement" in text:
            header_index = index
            break
    if header_index is None:
        return pd.DataFrame()

    headers = rows[header_index]
    records = []
    current_code, current_name = "", ""
    for row in rows[header_index + 1:]:
        if len(row) < len(headers):
            row += [""] * (len(headers) - len(row))
        values = dict(zip(headers, row[:len(headers)]))
        first_cell = str(row[0]).strip() if row else ""
        item_match = re.match(r"^\s*((?:GB|RB)[A-Z0-9]*)\s*[-–]\s*(.*)", first_cell, re.IGNORECASE)
        if item_match:
            current_code = item_match.group(1).upper()
            current_name = item_match.group(2).strip()
            continue
        if not str(values.get("Date", "")).strip():
            continue
        text = normalize(
            f"{values.get('Source', '')} {values.get('Reference', '')} "
            f"{values.get('Contact', '')} {values.get('Description', '')}"
        )
        if "امر تحميص مستودع البن الاخضر" in text:
            mtype = GREEN_OUT
        elif "تحميص وارد لمستودع البن المحمص" in text:
            mtype = ROASTED_IN
        elif "سند صرف قهوه محموصه فرع" in text:
            mtype = BRANCH_OUT
        elif "عينات القهوة المحمصة" in text:
            mtype = SAMPLE
        elif "receivable invoice" in text:
            mtype = SALE
        elif any(w in text for w in ["adjustment", "جرد", "revaluation", "تعديل"]):
            mtype = ADJUST
        else:
            mtype = OTHER
        coffee_type, family = item_meta(current_code)
        records.append({
            "year": year, "month": month, "mdate": values.get("Date", ""),
            "item_code": current_code, "item_name": current_name,
            "coffee_type": coffee_type, "family": family,
            "qty": number(values.get("QoH Movement", 0)),
            "value": number(values.get("Value Movement", 0)),
            "total": number(values.get("Total", 0)),
            "movement_type": mtype, "source": values.get("Source", ""),
            "reference": values.get("Reference", ""), "contact": values.get("Contact", ""),
            "description": values.get("Description", ""), "branch": "",
            "revenue_account": values.get("Revenue Account", ""),
            "inventory_account": values.get("Inventory Account", ""),
        })
    return pd.DataFrame(records)


def insert_financial(df, statement):
    if df.empty:
        return 0
    seq = [
        (statement, int(r["السنة"]), int(r["الشهر"]), str(r["كود الحساب"]),
         str(r["اسم الحساب"]), str(r["نوع الحساب"]), float(r["المبلغ"]))
        for _, r in df.iterrows()
    ]
    executemany(
        "INSERT INTO financial_lines(statement, year, month, account_code, account_name, account_type, amount) "
        "VALUES(?, ?, ?, ?, ?, ?, ?)",
        seq,
    )
    return len(seq)


def insert_movements(df, from_report=0):
    if df.empty:
        return 0
    cols = MOVE_COLS + ["from_report"]
    seq = []
    for _, r in df.iterrows():
        d = {c: r.get(c) for c in MOVE_COLS}
        d["created_at"] = now_str()
        d["created_by"] = current_user()
        d["from_report"] = int(from_report)
        seq.append(tuple(d.get(c) for c in cols))
    executemany(
        "INSERT INTO movements(" + ",".join(cols) + ") VALUES(" + ",".join(["?"] * len(cols)) + ")",
        seq,
    )
    # تحديث كتالوج الأصناف
    for code, name in df[["item_code", "item_name"]].drop_duplicates().values:
        upsert_item(code, name)
    return len(seq)


def upsert_balance(rec):
    """تخزين الرصيد الختامي المعتمد من التقرير (نقطة جديدة)."""
    code = (rec.get("item_code") or "").upper().strip()
    if not code:
        return
    coffee_type, family = item_meta(code)
    execute(
        "INSERT INTO item_balances(item_code, item_name, coffee_type, family, qty, value, "
        "year, month, source_file, updated_at) VALUES(?,?,?,?,?,?,?,?,?,?) "
        "ON CONFLICT(item_code) DO UPDATE SET item_name=excluded.item_name, "
        "coffee_type=excluded.coffee_type, family=excluded.family, qty=excluded.qty, "
        "value=excluded.value, year=excluded.year, month=excluded.month, "
        "source_file=excluded.source_file, updated_at=excluded.updated_at",
        (code, rec.get("item_name", ""), coffee_type, family,
         float(rec.get("qty") or 0), float(rec.get("value") or 0),
         rec.get("year"), rec.get("month"), rec.get("source_file", ""), now_str()),
    )
    upsert_item(code, rec.get("item_name", ""), coffee_type, family)


def insert_branch_sales(df, source_file=""):
    if df.empty:
        return 0
    seq = [
        (int(r["year"]), int(r["month"]), r["mdate"], r["invoice"], r["contact"],
         r["reference"], float(r["amount"]), r["status"], r["source"], r["branch"],
         source_file, now_str())
        for _, r in df.iterrows()
    ]
    executemany(
        "INSERT INTO branch_sales(year, month, mdate, invoice, contact, reference, "
        "amount, status, source, branch, source_file, imported_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", seq,
    )
    return len(seq)


def parse_inventory_xlsx(path, year, month):
    """يقرأ تقرير حركة الصنف (Xero) من ملف Excel ويعيد (حركات, أرصدة ختامية).
    الرصيد الختامي هو المرجع المعتمد (صف Closing Balance: الكمية من QoH والقيمة من Value Movement)."""
    try:
        raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
    except Exception:
        return pd.DataFrame(), pd.DataFrame()
    if raw.empty:
        return pd.DataFrame(), pd.DataFrame()
    header_idx = None
    for i in range(min(15, len(raw))):
        cells = [normalize(c) for c in raw.iloc[i].tolist()]
        if "date" in cells and any("qoh" in c for c in cells):
            header_idx = i
            break
    if header_idx is None:
        return pd.DataFrame(), pd.DataFrame()

    hdr = [normalize(c) for c in raw.iloc[header_idx].tolist()]

    def col_of(*keys):
        for j, h in enumerate(hdr):
            if any(k in h for k in keys):
                return j
        return None

    c_date = col_of("date") or 0
    c_source = col_of("source")
    c_ref = col_of("reference")
    c_contact = col_of("contact")
    c_desc = col_of("description")
    c_total = col_of("total")
    c_value = col_of("value movement")
    c_qoh = col_of("qoh")

    code_re = re.compile(r"^\s*((?:GB|RB)[A-Z0-9]*)\s*[-–]\s*(.*)", re.IGNORECASE)
    moves, balances = [], []
    cur_code, cur_name = "", ""

    def cell(row, idx):
        return row[idx] if idx is not None and idx < len(row) else ""

    for i in range(header_idx + 1, len(raw)):
        row = raw.iloc[i].tolist()
        c0 = str(cell(row, 0)).strip()
        low = normalize(c0)
        m = code_re.match(c0)
        if m:
            cur_code = m.group(1).upper()
            cur_name = m.group(2).strip()
            continue
        if low.startswith("total "):
            continue
        if low in ("opening balance", "closing balance") or low.startswith("closing bal"):
            if "closing" in low and cur_code:
                balances.append({
                    "item_code": cur_code, "item_name": cur_name,
                    "qty": number(cell(row, c_qoh)), "value": number(cell(row, c_value)),
                    "year": year, "month": month, "source_file": Path(path).name,
                })
            continue
        date_val = str(cell(row, c_date)).strip()
        if not date_val or date_val.lower() == "nan":
            continue
        if not cur_code:
            continue
        text = normalize(
            f"{cell(row, c_source)} {cell(row, c_ref)} {cell(row, c_contact)} {cell(row, c_desc)}"
        )
        if "امر تحميص مستودع البن الاخضر" in text:
            mtype = GREEN_OUT
        elif "تحميص وارد لمستودع البن المحمص" in text:
            mtype = ROASTED_IN
        elif "سند صرف قهوه محموصه فرع" in text:
            mtype = BRANCH_OUT
        elif "عينات القهوة المحمصة" in text:
            mtype = SAMPLE
        elif "receivable invoice" in text:
            mtype = SALE
        elif any(w in text for w in ["adjustment", "جرد", "revaluation", "تعديل"]):
            mtype = ADJUST
        else:
            mtype = OTHER
        coffee_type, family = item_meta(cur_code)
        moves.append({
            "year": year, "month": month, "mdate": date_val[:10],
            "item_code": cur_code, "item_name": cur_name,
            "coffee_type": coffee_type, "family": family,
            "qty": number(cell(row, c_qoh)), "value": number(cell(row, c_value)),
            "total": number(cell(row, c_total)),
            "movement_type": mtype, "source": str(cell(row, c_source)),
            "reference": str(cell(row, c_ref)), "contact": str(cell(row, c_contact)),
            "description": str(cell(row, c_desc)), "branch": "",
            "revenue_account": "", "inventory_account": "",
        })
    return pd.DataFrame(moves), pd.DataFrame(balances)


def parse_daily_sales(path, year, month):
    """يقرأ تقرير المبيعات اليومية لكل فرع (عمود Branches مقصود ويُوضَّح في التقرير)."""
    try:
        raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
    except Exception:
        return pd.DataFrame()
    if raw.empty:
        return pd.DataFrame()
    header_idx = None
    for i in range(min(20, len(raw))):
        cells = [normalize(c) for c in raw.iloc[i].tolist()]
        if "branches" in cells and any("invoice" in c for c in cells):
            header_idx = i
            break
    if header_idx is None:
        return pd.DataFrame()
    hdr = [normalize(c) for c in raw.iloc[header_idx].tolist()]

    def col_of(*keys):
        for j, h in enumerate(hdr):
            if any(k in h for k in keys):
                return j
        return None

    c_inv = col_of("invoice number") or 0
    c_contact = col_of("contact")
    c_date = col_of("invoice date")
    c_ref = col_of("reference")
    c_gross = col_of("gross", "amount", "total")
    c_status = col_of("status")
    c_source = col_of("source")
    c_branch = col_of("branches", "branch")
    rows = []
    for i in range(header_idx + 1, len(raw)):
        row = raw.iloc[i].tolist()
        inv = str(row[c_inv]).strip() if c_inv < len(row) else ""
        if not inv or inv.lower() == "nan" or normalize(inv).startswith("total"):
            continue
        d = str(row[c_date]).strip() if c_date is not None and c_date < len(row) else ""
        if not d or d.lower() == "nan":  # تخطّي صفوف الإجماليات (لا تاريخ لها)
            continue
        branch = str(row[c_branch]).strip() if c_branch is not None and c_branch < len(row) else ""
        if not branch or branch.lower() == "nan":
            branch = "غير محدد"
        mo = month
        mm = re.search(r"\d{4}-(\d{2})-\d{2}", d)
        if mm:
            mo = int(mm.group(1))
        rows.append({
            "year": year, "month": mo, "mdate": d[:10],
            "invoice": inv,
            "contact": str(row[c_contact]) if c_contact is not None and c_contact < len(row) else "",
            "reference": str(row[c_ref]) if c_ref is not None and c_ref < len(row) else "",
            "amount": number(row[c_gross]) if c_gross is not None and c_gross < len(row) else 0.0,
            "status": str(row[c_status]) if c_status is not None and c_status < len(row) else "",
            "source": str(row[c_source]) if c_source is not None and c_source < len(row) else "",
            "branch": branch,
        })
    return pd.DataFrame(rows)


MONTH_MAP = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def _pnl_category(section):
    s = normalize(section)
    if any(k in s for k in ["cost of", "direct cost", "تكلفة"]):
        return "cogs"
    if "other income" in s or "اخرى" in s and "ايراد" in s:
        return "other_income"
    if any(k in s for k in ["income", "revenue", "sales", "ايراد", "إيراد", "مبيعات", "دخل"]):
        return "revenue"
    if any(k in s for k in ["expense", "overhead", "depreciation", "admin", "operating",
                            "مصروف", "مصاريف", "اهلاك", "إهلاك", "تشغيل"]):
        return "expense"
    return None


def parse_pnl_monthly(path):
    """يقرأ قائمة الدخل الشهرية (أعمدة الأشهر Jan..Dec) إلى بنود شهرية نظيفة.
    يستبعد عمود Year to date ويصنّف البنود حسب القسم (إيراد/تكلفة مبيعات/مصروف)."""
    try:
        raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
    except Exception:
        return pd.DataFrame()
    if raw.empty:
        return pd.DataFrame()
    header_idx, month_cols = None, {}
    for i in range(min(20, len(raw))):
        cols = {}
        for j, c in enumerate(raw.iloc[i].tolist()):
            t = normalize(c)
            m = re.match(r"^([a-z]{3})[a-z]*\s*(20\d\d)", t)
            if m and m.group(1) in MONTH_MAP:
                cols[j] = (int(m.group(2)), MONTH_MAP[m.group(1)])
        if len(cols) >= 3:
            header_idx, month_cols = i, cols
            break
    if header_idx is None:
        return pd.DataFrame()

    code_re = re.compile(r"^\s*(\d+)\s*[-–]\s*(.*)$")
    section, records = "", []
    for i in range(header_idx + 1, len(raw)):
        row = raw.iloc[i].tolist()
        c0 = str(row[0]).strip() if row else ""
        if not c0 or c0.lower() == "nan":
            continue
        low = normalize(c0)
        if low.startswith(("total", "gross profit", "net profit", "إجمالي", "صافي", "مجمل")):
            continue
        m = code_re.match(c0)
        has_values = any(number(row[j]) for j in month_cols if j < len(row))
        if not m and not has_values:
            section = c0  # عنوان قسم (Trading Income / Cost of Sales / Operating Expenses ...)
            continue
        category = _pnl_category(section)
        if category is None:
            continue
        code = m.group(1) if m else ""
        name = m.group(2).strip() if m else c0
        for j, (yr, mo) in month_cols.items():
            if j >= len(row):
                continue
            amt = number(row[j])
            if amt:
                records.append({
                    "year": yr, "month": mo, "account_code": code,
                    "account_name": name, "category": category, "amount": amt,
                })
    return pd.DataFrame(records)


def parse_balance_sheet(path):
    """يجمع بنود الميزانية يدويًا (صفوف Total تُصدَّر أصفارًا) لكل سنة:
    الأصول المتداولة، المخزون، الخصوم المتداولة، إجمالي الأصول، حقوق الملكية."""
    try:
        raw = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
    except Exception:
        return pd.DataFrame()
    if raw.empty:
        return pd.DataFrame()
    header_idx, year_cols = None, {}
    for i in range(min(12, len(raw))):
        cols = {}
        for j, c in enumerate(raw.iloc[i].tolist()):
            mm = re.search(r"(20\d\d)", str(c))
            if mm and normalize(c) not in ("account",):
                cols[j] = int(mm.group(1))
        if cols and any("account" in normalize(x) for x in raw.iloc[i].tolist()):
            header_idx, year_cols = i, cols
            break
    if header_idx is None:
        return pd.DataFrame()

    agg = {y: dict(current_assets=0.0, inventory=0.0, current_liabilities=0.0,
                   total_assets=0.0, total_equity=0.0) for y in year_cols.values()}
    top, sub = "", ""
    code_re = re.compile(r"^\s*\d+\s*[-–]")
    for i in range(header_idx + 1, len(raw)):
        row = raw.iloc[i].tolist()
        c0 = str(row[0]).strip() if len(row) > 0 else ""
        c1 = str(row[1]).strip() if len(row) > 1 else ""
        label = c1 if c1 and c1.lower() != "nan" else c0
        low = normalize(label)
        vals = {y: number(row[j]) for j, y in year_cols.items() if j < len(row)}
        has_val = any(v for v in vals.values())
        is_detail = bool(code_re.match(c1)) or (has_val and not low.startswith("total"))
        if c0 and c0.lower() != "nan" and not code_re.match(c0):
            top = normalize(c0)
        if not is_detail:
            if low.startswith("total") or not label or label.lower() == "nan":
                continue
            sub = low  # عنوان قسم فرعي (Bank / Current Assets / Current Liabilities ...)
            continue
        is_current_asset = ("current asset" in sub) or ("bank" in sub) or \
                           ("asset" in top and "current" in sub)
        is_current_liab = ("current liabilit" in sub) or ("current" in sub and "liab" in top)
        is_inventory = ("مخزون" in low) or ("inventory" in low)
        for y, v in vals.items():
            if "asset" in top:
                agg[y]["total_assets"] += v
            if "equity" in top:
                agg[y]["total_equity"] += v
            if is_current_asset:
                agg[y]["current_assets"] += v
            if is_inventory:
                agg[y]["inventory"] += v
            if is_current_liab:
                agg[y]["current_liabilities"] += v
    rows = [dict(year=y, source_file=Path(path).name, **d) for y, d in agg.items()]
    return pd.DataFrame(rows)


def insert_pnl_monthly(df, source_file=""):
    if df.empty:
        return 0
    # إبطال التكرار على مستوى (سنة/شهر): يسمح بإضافة ملف كل شهر على حِدة
    # دون أن يمحو ملفُ شهرٍ بياناتِ الأشهر الأخرى، ودون مضاعفة عند إعادة الاستيراد.
    pairs = sorted(set((int(r["year"]), int(r["month"])) for _, r in df.iterrows()))
    for y, mo in pairs:
        execute("DELETE FROM pnl_monthly WHERE year=? AND month=?", (y, mo))
    seq = [
        (int(r["year"]), int(r["month"]), r["account_code"], r["account_name"],
         r["category"], float(r["amount"]), source_file, now_str())
        for _, r in df.iterrows()
    ]
    executemany(
        "INSERT INTO pnl_monthly(year, month, account_code, account_name, category, "
        "amount, source_file, imported_at) VALUES(?,?,?,?,?,?,?,?)", seq,
    )
    return len(seq)


def upsert_fin_figures(df):
    if df.empty:
        return 0
    n = 0
    for _, r in df.iterrows():
        execute(
            "INSERT INTO fin_figures(year, current_assets, inventory, current_liabilities, "
            "total_assets, total_equity, source_file, updated_at) VALUES(?,?,?,?,?,?,?,?) "
            "ON CONFLICT(year) DO UPDATE SET current_assets=excluded.current_assets, "
            "inventory=excluded.inventory, current_liabilities=excluded.current_liabilities, "
            "total_assets=excluded.total_assets, total_equity=excluded.total_equity, "
            "source_file=excluded.source_file, updated_at=excluded.updated_at",
            (int(r["year"]), float(r["current_assets"]), float(r["inventory"]),
             float(r["current_liabilities"]), float(r["total_assets"]),
             float(r["total_equity"]), r.get("source_file", ""), now_str()),
        )
        n += 1
    return n


def save_raw_statement(statement, year, month, df):
    payload = df.to_json(orient="split", force_ascii=False)
    execute(
        "INSERT INTO raw_statements(statement, year, month, payload, imported_at) VALUES(?, ?, ?, ?, ?)",
        (statement, year, month, payload, now_str()),
    )


def import_inventory_report(log):
    """يستورد تقارير حركة الصنف (Excel) من المجلد الرئيسي ومن مجلدات الأشهر (YYYY-MM):
    حركات كل الأعوام/الأشهر + الرصيد الختامي المعتمد من أحدث فترة فقط."""
    by_period = {}
    for path in _scan_report_xlsx():
        nm = normalize(path.name)
        if not (("inventory" in nm and "detail" in nm) or ("تفصيل" in nm and "inventory" in nm)):
            continue
        if "summary" in nm or "مجمع" in nm:
            continue
        pm = re.match(r"(20\d\d)[-_](\d{1,2})", path.parent.name)
        if pm:
            year, month = int(pm.group(1)), int(pm.group(2))
        else:
            ym = re.search(r"(20\d\d)", path.name)
            year, month = (int(ym.group(1)) if ym else 2025), 12
        # ملف واحد لكل فترة، مع تفضيل النسخة داخل مجلد الإدخال الشهري كمرجع معتمد
        key = (year, month)
        prev = by_period.get(key)
        in_input = INPUT_DIR in path.parents
        if prev is None or (in_input and INPUT_DIR not in prev.parents):
            by_period[key] = path
    if not by_period:
        return
    latest = max(by_period)  # أحدث تاريخ (سنة، شهر) في ملفات الإدخال = مرجع الرصيد الحالي
    for (year, month) in sorted(by_period):
        path = by_period[(year, month)]
        moves, balances = parse_inventory_xlsx(path, year, month)
        n = insert_movements(moves, from_report=1)
        b = 0
        if (year, month) == latest:
            for _, r in balances.iterrows():
                upsert_balance(r.to_dict())
                b += 1
        log.append(f"{path.name} ← حركات {year}-{month:02d}: {n} | أرصدة ختامية معتمدة: {b}")


def import_daily_sales(log):
    """يستورد تقرير المبيعات اليومية لكل فرع (عمود Branches) من المجلد الرئيسي
    ومن مجلدات الأشهر (YYYY-MM) — لكل الأعوام، بدون تكرار الفترة الواحدة."""
    by_period = {}
    for path in _scan_report_xlsx():
        nm = normalize(path.name)
        if not ("المبيعات اليومية" in nm or ("daily" in nm and "sale" in nm) or "لكل فرع" in nm):
            continue
        pm = re.match(r"(20\d\d)[-_](\d{1,2})", path.parent.name)
        if pm:
            year, month = int(pm.group(1)), int(pm.group(2))
        else:
            ym = re.search(r"(20\d\d)", path.name)
            year, month = (int(ym.group(1)) if ym else 2025), 12
        by_period.setdefault((year, month), path)
    for (year, month) in sorted(by_period):
        path = by_period[(year, month)]
        df = parse_daily_sales(path, year, month)
        n = insert_branch_sales(df, source_file=path.name)
        log.append(f"{path.name} ← فواتير المبيعات حسب الفرع {year}: {n}")


def _scan_report_xlsx():
    """كل ملفات Excel في المجلد الرئيسي + مجلدات الإدخال الشهرية (بدون الأرشيف)."""
    seen = set()
    paths = list(BASE_DIR.glob("*.xlsx"))
    if INPUT_DIR.exists():
        paths += list(INPUT_DIR.rglob("*.xlsx"))
    for p in paths:
        rp = str(p.resolve())
        if rp in seen:
            continue
        seen.add(rp)
        yield p


def import_pnl_monthly(log):
    """قائمة الدخل الشهرية (أعمدة الأشهر) → جدول تحليلي نظيف لكل شهر."""
    for path in _scan_report_xlsx():
        nm = normalize(path.name)
        if "profit and loss" not in nm and "الدخل" not in nm and "by_month" not in nm \
           and "by month" not in nm and "financial_year_by_month" not in nm \
           and "current_financial_year" not in nm and "current financial year" not in nm:
            continue
        df = parse_pnl_monthly(path)
        if df.empty:
            continue
        n = insert_pnl_monthly(df, source_file=path.name)
        yrs = sorted(set(int(y) for y in df["year"].unique()))
        log.append(f"{path.name} ← بنود قائمة الدخل الشهرية: {n} (سنوات: {', '.join(map(str, yrs))})")


def import_balance_sheet_figures(log):
    """الميزانية العمومية → أصول/خصوم متداولة ومخزون لكل سنة (لحساب النسب)."""
    for path in _scan_report_xlsx():
        nm = normalize(path.name)
        if "balance sheet" not in nm and "الميزانية" not in nm and "balance_sheet" not in nm:
            continue
        df = parse_balance_sheet(path)
        if df.empty:
            continue
        n = upsert_fin_figures(df)
        log.append(f"{path.name} ← أرقام الميزانية للنسب المالية: {n} سنة")


def import_from_folders():
    """يعيد استخدام منطق السكربت الأصلي لاستيراد كل المجلدات الشهرية إلى قاعدة البيانات."""
    log = []
    # إعادة استيراد غير مكرّرة: نحذف بيانات التقرير السابقة (وأي حركات قديمة بدون وسم)
    # مع الحفاظ على الإدخالات اليدوية (from_report=0 صراحةً)
    execute("DELETE FROM movements WHERE from_report=1 OR from_report IS NULL")
    execute("DELETE FROM item_balances")
    execute("DELETE FROM branch_sales")
    execute("DELETE FROM pnl_monthly")

    chart, chart_name = load_chart_of_accounts()
    if not chart.empty:
        for _, r in chart.iterrows():
            upsert_account(r["كود الحساب"], r["اسم الحساب"], r["نوع الحساب"])
        log.append(f"شجرة الحسابات: {len(chart)} حساب ({chart_name})")
    else:
        log.append("تحذير: تعذّر قراءة شجرة الحسابات.")

    # النقاط الجديدة: تقرير المخزون المعتمد + المبيعات حسب الفرع
    import_inventory_report(log)
    import_daily_sales(log)
    import_pnl_monthly(log)
    import_balance_sheet_figures(log)

    folders = sorted([f for f in INPUT_DIR.iterdir() if f.is_dir()]) if INPUT_DIR.exists() else []
    if not folders:
        log.append("ملاحظة: لا توجد مجلدات أشهر إضافية داخل مجلد الإدخال (تم الاعتماد على تقارير المجلد الرئيسي).")
        audit("import_folders", " | ".join(log))
        return log

    for folder in folders:
        year, month = extract_year_month(folder)
        for path in folder.iterdir():
            ext = path.suffix.lower()
            if ext == ".csv":
                inv = parse_inventory(path, year, month)
                n = insert_movements(inv, from_report=1)
                log.append(f"{path.name} ← حركة مخزون: {n}")
            elif ext in (".xlsx", ".xls"):
                data = read_excel(path)
                if data.empty:
                    continue
                fname = path.name.lower()
                sig = " ".join(normalize(c) for c in data.columns)
                if "profit and loss" in fname or "profit and loss" in sig:
                    parsed = parse_profit_and_loss(data, year, chart)
                    n = insert_financial(parsed, "pnl")
                    log.append(f"{path.name} ← أرباح وخسائر: {n}")
                elif "trial balance" in fname or "trial balance" in sig:
                    save_raw_statement("trial", year, month, data)
                    log.append(f"{path.name} ← ميزان مراجعة (حُفظ خام): {len(data)}")
                elif "balance sheet" in fname or "balance sheet" in sig:
                    save_raw_statement("balance", year, month, data)
                    log.append(f"{path.name} ← ميزانية عمومية (حُفظت خام): {len(data)}")
    audit("import_folders", " | ".join(log))
    return log


# ==========================================================================
# التصدير إلى Excel (منقول من السكربت الأصلي)
# ==========================================================================
def write_sheet(writer, name, data):
    if data is None or data.empty:
        data = pd.DataFrame({"ملاحظة": ["لا توجد بيانات متاحة لهذه الصفحة"]})
    data.to_excel(writer, sheet_name=name[:31], index=False)


def style_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    for ws in workbook.worksheets:
        ws.sheet_view.rightToLeft = True
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
    workbook.save(path)


def build_exports():
    settings = get_settings()
    accounts = list_accounts().rename(columns={"code": "كود الحساب", "name": "اسم الحساب", "type": "نوع الحساب"})
    pnl = pnl_df()
    inv = movements_df()
    budget = query_df(
        "SELECT year AS 'السنة', month AS 'الشهر', account_code AS 'كود الحساب', "
        "account_name AS 'اسم الحساب', account_type AS 'نوع الحساب', actual_prev AS 'فعلي الأساس', "
        "growth AS 'نمو', proposed AS 'موازنة مقترحة', approved AS 'موازنة معتمدة' FROM budget"
    )
    forecast = build_forecast(pnl)
    inv_summary = build_inventory_summary(inv)
    roasting = build_roasting(inv, s_float(settings, "roast_yield"), s_float(settings, "roast_tolerance"))
    stock = current_stock()
    branch_detail, branch_by = branch_sales_summary()

    with pd.ExcelWriter(DASHBOARD_PATH, engine="openpyxl") as writer:
        write_sheet(writer, "شجرة الحسابات", accounts)
        write_sheet(writer, "قائمة الأرباح والخسائر", pnl)
        write_sheet(writer, "الموازنة", budget)
        write_sheet(writer, "التوقعات", forecast)
        write_sheet(writer, "الرصيد الحالي المعتمد", stock)
        write_sheet(writer, "حركة المخزون", inv)
        write_sheet(writer, "تحليل المخزون", inv_summary)
        write_sheet(writer, "تحليل التحميص", roasting)
        write_sheet(writer, "المبيعات حسب الفرع", branch_by)
        write_sheet(writer, "تفاصيل فواتير الفروع", branch_detail)
    style_workbook(DASHBOARD_PATH)

    alerts_review = roasting[roasting["الحالة"] == "يحتاج مراجعة"] if not roasting.empty else pd.DataFrame()
    with pd.ExcelWriter(ALERTS_PATH, engine="openpyxl") as writer:
        write_sheet(writer, "فروقات التحميص", alerts_review)
        write_sheet(writer, "ملخص المخزون", inv_summary)
    style_workbook(ALERTS_PATH)
    audit("export_excel", DASHBOARD_PATH.name)
    return DASHBOARD_PATH, ALERTS_PATH


# ==========================================================================
# طبقة التوافق مع إصدارات Streamlit المختلفة
# ==========================================================================
def rerun():
    fn = getattr(st, "rerun", None) or getattr(st, "experimental_rerun", None)
    if fn:
        fn()


def editable_table(df, key):
    if hasattr(st, "data_editor"):
        return st.data_editor(df, use_container_width=True, key=key, num_rows="fixed")
    if hasattr(st, "experimental_data_editor"):
        return st.experimental_data_editor(df, use_container_width=True, key=key)
    st.warning("⚠️ إصدار Streamlit لديك قديم لا يدعم الجداول القابلة للتحرير. "
               "يرجى التحديث بالأمر: pip install -U streamlit")
    st.dataframe(df, use_container_width=True)
    return None


# ==========================================================================
# واجهة المستخدم
# ==========================================================================
def setup_page():
    st.set_page_config(page_title="داشبورد الكوب الأول", page_icon="☕", layout="wide")
    st.markdown(
        """
        <style>
        .stApp, .main, section[data-testid="stSidebar"] { direction: rtl; }
        h1, h2, h3, h4, h5, p, label, .stMarkdown { text-align: right; }
        .stDataFrame, .stTable { direction: rtl; }
        div[data-testid="stMetricValue"] { direction: ltr; }
        </style>
        """,
        unsafe_allow_html=True,
    )


def sidebar():
    st.sidebar.title("☕ الكوب الأول")
    st.sidebar.caption("نظام إدارة المخزون والمبيعات التفاعلي")
    st.session_state["user_name"] = st.sidebar.text_input("اسم المستخدم", value=current_user())
    pages = [
        "🏠 لوحة المعلومات",
        "🧾 المبيعات اليومية",
        "🛵 مبيعات التطبيقات والموقع",
        "🔥 دفعات التحميص",
        "📦 المخزون والحركات",
        "💰 الموازنة والتوقعات",
        "📊 التحليل المالي",
        "📥 استيراد البيانات",
        "📚 الأصناف والحسابات",
        "🔔 التنبيهات",
        "⚙️ الإعدادات",
        "📤 تصدير Excel",
    ]
    count = open_alerts_count()
    if count:
        st.sidebar.error(f"🔔 لديك {count} تنبيه غير محلول")
    return st.sidebar.radio("انتقل إلى", pages)


def item_selector(label, prefix=None, key=None):
    """قائمة اختيار الأصناف مع البحث. prefix=GB للأخضر، RB للمحمص."""
    items = list_items()
    if prefix:
        items = items[items["code"].str.upper().str.startswith(prefix)]
    if items.empty:
        st.info("لا توجد أصناف بعد. أضِف أصنافًا أو استورد بياناتك أولًا.")
        return None
    options = items["code"].tolist()
    labels = {r["code"]: f"{r['code']} — {r['name']}" for _, r in items.iterrows()}
    return st.selectbox(label, options, format_func=lambda c: labels.get(c, c), key=key)


# ---- صفحة لوحة المعلومات ----
def page_dashboard():
    st.header("🏠 لوحة المعلومات")
    inv = movements_df()
    stock = current_stock()
    if inv.empty:
        st.info("لا توجد بيانات بعد. ابدأ من صفحة «استيراد البيانات» أو أدخل مبيعات ودفعات تحميص.")
        return

    sales = inv[inv["نوع الحركة"] == SALE]
    total_sales = sales["الإجمالي"].sum()
    stock_value = stock["قيمة الرصيد"].sum() if not stock.empty else 0
    settings = get_settings()
    roasting = build_roasting(inv, s_float(settings, "roast_yield"), s_float(settings, "roast_tolerance"))
    avg_yield = 0
    if not roasting.empty and roasting["كمية الأخضر المحولة"].sum():
        avg_yield = roasting["كمية المحمص الناتجة"].sum() / roasting["كمية الأخضر المحولة"].sum() * 100

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("إجمالي المبيعات", f"{total_sales:,.0f}")
    c2.metric("عدد الأصناف", f"{len(stock)}")
    c3.metric("قيمة المخزون", f"{stock_value:,.0f}")
    c4.metric("متوسط كفاءة التحميص", f"{avg_yield:,.1f}%")

    st.divider()
    col_a, col_b = st.columns(2)
    with col_a:
        st.subheader("المبيعات حسب الشهر")
        if not sales.empty:
            by_month = sales.groupby("الشهر")["الإجمالي"].sum()
            st.bar_chart(by_month)
        else:
            st.caption("لا توجد مبيعات مسجلة بعد.")
    with col_b:
        st.subheader("توزيع الحركات حسب النوع")
        by_type = inv.groupby("نوع الحركة")["حركة الكمية"].count()
        st.bar_chart(by_type)

    st.subheader("الرصيد الحالي للأصناف")
    st.caption("تُعرض الأصناف التي رصيدها أكبر من صفر فقط.")
    positive = stock[stock["الرصيد الحالي"] > 0] if not stock.empty else stock
    st.dataframe(positive, use_container_width=True, hide_index=True)

    # النقطة الجديدة: المبيعات حسب الفرع (عمود Branches من تقرير المبيعات اليومية)
    detail, by_branch = branch_sales_summary()
    if not by_branch.empty:
        st.divider()
        st.subheader("المبيعات حسب الفرع")
        total_branch = detail["المبلغ"].sum()
        st.caption(f"عدد الفواتير: {len(detail):,} — إجمالي المبيعات: {total_branch:,.0f}")
        cc1, cc2 = st.columns([1, 1])
        with cc1:
            st.dataframe(by_branch, use_container_width=True, hide_index=True)
        with cc2:
            st.bar_chart(by_branch.set_index("الفرع")["إجمالي المبيعات"])
        with st.expander("تفاصيل الفواتير لكل فرع"):
            branches = ["(الكل)"] + by_branch["الفرع"].tolist()
            pick = st.selectbox("اختر الفرع", branches, key="dash_branch_pick")
            view = detail if pick == "(الكل)" else detail[detail["الفرع"] == pick]
            st.dataframe(view, use_container_width=True, hide_index=True)


# ---- صفحة المبيعات اليومية ----
def page_sales():
    st.header("🧾 المبيعات اليومية")
    st.caption("سجّل عملية بيع؛ يُخصم المخزون تلقائيًا ويُحدَّث لوح المعلومات فورًا.")
    with st.form("sale_form", clear_on_submit=True):
        col1, col2 = st.columns(2)
        with col1:
            d = st.date_input("التاريخ", value=date.today())
            item = item_selector("الصنف", key="sale_item")
            qty = st.number_input("الكمية", min_value=0.0, step=1.0)
        with col2:
            price = st.number_input("سعر الوحدة", min_value=0.0, step=1.0)
            branch = st.text_input("الفرع")
            customer = st.text_input("العميل")
        reference = st.text_input("رقم المرجع / الفاتورة")
        submitted = st.form_submit_button("💾 حفظ عملية البيع")

    if submitted:
        if not item or qty <= 0:
            st.error("اختر صنفًا وأدخل كمية أكبر من صفر.")
            return
        available = stock_of(item)
        add_movement(
            year=d.year, month=d.month, mdate=d.isoformat(), item_code=item,
            item_name=_name_of(item), qty=-abs(qty), total=qty * price, value=0,
            movement_type=SALE, branch=branch, contact=customer, reference=reference,
            description="بيع يدوي",
        )
        audit("sale", f"{item} x{qty} @ {price}")
        st.success(f"تم تسجيل بيع {qty:g} من {item} بقيمة {qty * price:,.0f}.")
        if available - qty < 0:
            add_alert("تحذير", "مخزون", f"الصنف {item} أصبح رصيده سالبًا بعد البيع ({available - qty:g}).")
            st.warning(f"⚠️ تنبيه: الرصيد أصبح سالبًا ({available - qty:g}). سُجّل تنبيه.")

    _sales_summary_section()


def _sales_summary_section():
    """ملخص مختصر ومجمّع لمبيعات العملاء والفروع من تقرير الإكسيل — أساس سنوي/شهري."""
    st.divider()
    st.subheader("📋 ملخص مبيعات العملاء والفروع (من تقرير الإكسيل)")
    syears = query_df("SELECT DISTINCT year FROM branch_sales ORDER BY year")
    if syears.empty:
        st.info("لا توجد بيانات مبيعات مستوردة بعد. استورد «تقرير المبيعات اليومية لكل فرع» من صفحة الاستيراد.")
        return
    yrs = [int(y) for y in syears["year"].tolist()]
    c1, c2 = st.columns([1, 2])
    basis = c1.radio("الأساس", ["سنوي", "شهري"], horizontal=True, key="bs_basis")
    year = c2.selectbox("السنة", yrs, index=len(yrs) - 1, key="bs_year")
    month = None
    if basis == "شهري":
        mons = query_df("SELECT DISTINCT month FROM branch_sales WHERE year=? ORDER BY month", (int(year),))
        mlist = [int(x) for x in mons["month"].tolist()] or list(range(1, 13))
        month = st.selectbox("الشهر", mlist,
                             format_func=lambda x: f"{x} - {MONTH_NAMES_AR[x - 1]}", key="bs_month")

    g = branch_sales_grouped(year, month)
    if g.empty:
        st.info("لا توجد فواتير في هذه الفترة.")
        return
    label = f"عام {int(year)}" if not month else f"{MONTH_NAMES_AR[month - 1]} {int(year)}"
    k1, k2, k3, k4 = st.columns(4)
    k1.metric("إجمالي المبيعات", f"{g['إجمالي المبيعات'].sum():,.0f}")
    k2.metric("عدد الفواتير", f"{int(g['عدد الفواتير'].sum()):,}")
    k3.metric("عدد العملاء", f"{g['العميل'].nunique():,}")
    k4.metric("عدد الفروع", f"{g['الفرع'].nunique():,}")
    st.caption(f"مبيعات {label} — مجمّعة حسب العميل (Contact) مع توضيح الفرع (Branches).")
    per_cust = g.groupby("العميل", as_index=False).agg(
        **{"عدد الفواتير": ("عدد الفواتير", "sum"), "مجموع المبيعات": ("إجمالي المبيعات", "sum")}
    )
    brs = (g.groupby("العميل")["الفرع"]
           .apply(lambda s: "، ".join(sorted(set(str(x) for x in s))))
           .reset_index().rename(columns={"الفرع": "الفروع"}))
    per_cust = (per_cust.merge(brs, on="العميل", how="left")
                .sort_values("مجموع المبيعات", ascending=False).reset_index(drop=True))
    per_cust["مجموع المبيعات"] = per_cust["مجموع المبيعات"].round(2)
    st.dataframe(per_cust, use_container_width=True, hide_index=True)

    with st.expander("تفصيل حسب العميل والفرع"):
        detail = g.copy()
        detail["إجمالي المبيعات"] = detail["إجمالي المبيعات"].round(2)
        st.dataframe(detail, use_container_width=True, hide_index=True)

    by_branch = g.groupby("الفرع", as_index=False).agg(
        **{"عدد الفواتير": ("عدد الفواتير", "sum"), "إجمالي المبيعات": ("إجمالي المبيعات", "sum")}
    ).sort_values("إجمالي المبيعات", ascending=False).reset_index(drop=True)
    by_branch["إجمالي المبيعات"] = by_branch["إجمالي المبيعات"].round(2)
    with st.expander("عرض مختصر حسب الفرع"):
        st.dataframe(by_branch, use_container_width=True, hide_index=True)


def _name_of(code):
    items = list_items()
    row = items[items["code"] == code]
    return row["name"].iloc[0] if not row.empty else code


# ---- صفحة دفعات التحميص ----
def page_roasting():
    st.header("🔥 دفعات التحميص")
    st.caption("أدخل كمية البن الأخضر الخارجة والمحمص الناتج؛ تُحسب نسبة الإنتاج (Yield) لحظيًا مع تنبيه عند تجاوز الحد.")
    settings = get_settings()
    target = s_float(settings, "roast_yield")
    tol = s_float(settings, "roast_tolerance")

    green = st.number_input("كمية البن الأخضر الخارجة (كجم)", min_value=0.0, step=1.0, key="g_kg")
    roasted = st.number_input("كمية البن المحمص الناتجة (كجم)", min_value=0.0, step=1.0, key="r_kg")

    if green > 0:
        yield_ratio = roasted / green
        diff = yield_ratio - target
        ok = abs(diff) <= tol
        c1, c2, c3 = st.columns(3)
        c1.metric("النسبة الفعلية", f"{yield_ratio * 100:,.1f}%")
        c2.metric("النسبة المستهدفة", f"{target * 100:,.1f}%")
        c3.metric("الحالة", "✅ مقبول" if ok else "⚠️ يحتاج مراجعة")
        if ok:
            st.success("النسبة ضمن الحد المسموح.")
        else:
            st.warning(f"النسبة خارج الحد المسموح (±{tol * 100:.1f}%).")

    with st.form("roast_form", clear_on_submit=True):
        col1, col2 = st.columns(2)
        with col1:
            d = st.date_input("تاريخ الدفعة", value=date.today())
            green_item = item_selector("صنف البن الأخضر", prefix="GB", key="roast_green")
        with col2:
            reference = st.text_input("رقم الدفعة / المرجع")
            roasted_item = item_selector("صنف البن المحمص", prefix="RB", key="roast_roasted")
        submitted = st.form_submit_button("💾 حفظ دفعة التحميص")

    if submitted:
        if not green_item or not roasted_item or green <= 0 or roasted <= 0:
            st.error("اختر صنفي الأخضر والمحمص وأدخل كميتين أكبر من صفر.")
            return
        add_movement(
            year=d.year, month=d.month, mdate=d.isoformat(), item_code=green_item,
            item_name=_name_of(green_item), qty=-abs(green), movement_type=GREEN_OUT,
            reference=reference, description="دفعة تحميص",
        )
        add_movement(
            year=d.year, month=d.month, mdate=d.isoformat(), item_code=roasted_item,
            item_name=_name_of(roasted_item), qty=abs(roasted), movement_type=ROASTED_IN,
            reference=reference, description="دفعة تحميص",
        )
        ratio = roasted / green if green else 0
        audit("roast_batch", f"{green_item}->{roasted_item} yield={ratio:.3f}")
        st.success(f"تم حفظ الدفعة. النسبة الفعلية {ratio * 100:.1f}%.")
        if abs(ratio - target) > tol:
            add_alert("مراجعة", "تحميص",
                      f"دفعة {reference or '-'}: نسبة الإنتاج {ratio * 100:.1f}% خارج الحد "
                      f"(المستهدف {target * 100:.1f}% ±{tol * 100:.1f}%).")
            st.warning("⚠️ سُجّل تنبيه لأن النسبة خارج الحد المسموح.")


# ---- صفحة المخزون والحركات ----
def _render_coffee_type_totals(df, qty_col, val_col, qty_label, val_label, caption):
    """تجميع الكمية والمبلغ حسب نوع البن (أخضر / محمص) للفترة المعروضة."""
    if df is None or df.empty or "نوع البن" not in df.columns:
        return
    g = df.copy()
    g["نوع البن"] = g["نوع البن"].apply(
        lambda x: "🟢 بن أخضر" if "أخضر" in str(x) else ("🟤 بن محمص" if "محمص" in str(x) else "أخرى"))
    agg = (g.groupby("نوع البن")
             .agg(**{qty_label: (qty_col, "sum"), val_label: (val_col, "sum")})
             .reset_index())
    order = {"🟢 بن أخضر": 0, "🟤 بن محمص": 1, "أخرى": 2}
    agg["_o"] = agg["نوع البن"].map(order).fillna(9)
    agg = agg.sort_values("_o").drop(columns="_o")
    total = pd.DataFrame({"نوع البن": ["الإجمالي"],
                          qty_label: [agg[qty_label].sum()], val_label: [agg[val_label].sum()]})
    out = pd.concat([agg, total], ignore_index=True)
    out[qty_label] = out[qty_label].round(1)
    out[val_label] = out[val_label].round(1)
    st.markdown(f"**{caption} (كميات ومبلغ):**")
    st.dataframe(out, use_container_width=True, hide_index=True)


def _render_movement_summary(summ, caption, key_all):
    if summ is None or summ.empty:
        st.info("لا توجد حركات لهذه الفترة.")
        return
    st.caption(caption)
    m1, m2, m3 = st.columns(3)
    m1.metric("إجمالي الوارد", f"{summ['الوارد'].sum():,.0f}")
    m2.metric("إجمالي المنصرف", f"{summ['المنصرف'].sum():,.0f}")
    m3.metric("صافي الحركة", f"{summ['صافي الحركة'].sum():,.0f}")
    _render_coffee_type_totals(summ, "صافي الحركة", "قيمة الحركة",
                               "صافي الحركة", "قيمة الحركة", "تجميع الحركة حسب نوع البن")
    show_all = st.checkbox("إظهار الأصناف بلا حركة أيضًا", value=False, key=key_all)
    view = summ if show_all else summ[(summ["الوارد"] != 0) | (summ["المنصرف"] != 0)]
    st.dataframe(view, use_container_width=True, hide_index=True)


def page_inventory():
    st.header("📦 المخزون والحركات")
    tab1, tab_m, tab_y, tab2, tab3 = st.tabs(
        ["الرصيد الحالي", "🗓️ رصيد شهري", "📅 رصيد سنوي", "📋 بيان الصنف المجمّع", "سجل الحركات"]
    )

    with tab1:
        stock = current_stock()
        if stock.empty:
            st.info("لا توجد حركات مخزون بعد.")
        else:
            show_all = st.checkbox("إظهار كل الأصناف (بما فيها الصفر والسالب)", value=False)
            view = stock if show_all else stock[stock["الرصيد الحالي"] > 0]
            st.caption("افتراضيًا تُعرض الأصناف التي رصيدها أكبر من صفر فقط.")
            _render_coffee_type_totals(view, "الرصيد الحالي", "قيمة الرصيد",
                                       "الكمية", "القيمة", "تجميع رصيد المخزون حسب نوع البن")
            st.dataframe(view, use_container_width=True, hide_index=True)
            neg = stock[stock["الرصيد الحالي"] < 0]
            if not neg.empty:
                st.warning(f"⚠️ يوجد {len(neg)} صنف برصيد سالب (فعّل الخيار أعلاه لعرضها).")

    with tab_m:
        yrs = movement_years()
        if not yrs:
            st.info("لا توجد حركات مخزون بعد.")
        else:
            c1, c2 = st.columns(2)
            y = c1.selectbox("السنة", yrs, index=len(yrs) - 1, key="inv_m_year")
            mons = movement_months(y)
            if not mons:
                st.info("لا توجد حركات لهذه السنة.")
            else:
                mo = c2.selectbox("الشهر", mons,
                                  format_func=lambda x: f"{x} - {MONTH_NAMES_AR[x - 1]}", key="inv_m_month")
                _render_movement_summary(
                    movement_period_summary(y, mo),
                    f"حركة المخزون لكل صنف خلال {MONTH_NAMES_AR[mo - 1]} {y} (من مدخلات الإكسيل).",
                    "inv_m_all",
                )

    with tab_y:
        yrs = movement_years()
        if not yrs:
            st.info("لا توجد حركات مخزون بعد.")
        else:
            y = st.selectbox("السنة", yrs, index=len(yrs) - 1, key="inv_y_year")
            _render_movement_summary(
                movement_period_summary(y, None),
                f"إجمالي حركة المخزون لكل صنف خلال عام {y} (من مدخلات الإكسيل).",
                "inv_y_all",
            )

    with tab2:
        st.subheader("📋 بيان مجمّع للصنف — من سجل الحركات")
        st.caption("لكل صنف: ماذا تم عليه من تحميص، فروقات جردية، بيع، شراء/وارد، صرف للفروع، عينات — "
                   "مفلتراً على مستوى العام والشهر.")
        yrs = movement_years()
        if not yrs:
            st.info("لا توجد حركات مخزون بعد.")
        else:
            c1, c2 = st.columns(2)
            basis = c1.radio("الأساس", ["سنوي", "شهري"], horizontal=True, key="stmt_basis")
            y = c2.selectbox("السنة", yrs, index=len(yrs) - 1, key="stmt_year")
            mo = None
            if basis == "شهري":
                mons = movement_months(y)
                if mons:
                    mo = st.selectbox("الشهر", mons,
                                      format_func=lambda x: f"{x} - {MONTH_NAMES_AR[x - 1]}", key="stmt_month")
            stmt = item_movement_statement(y, mo)
            if stmt.empty:
                st.info("لا توجد حركات في هذه الفترة.")
            else:
                label = f"عام {y}" if not mo else f"{MONTH_NAMES_AR[mo - 1]} {y}"
                totals = {c: stmt[c].sum() for c in stmt.columns if c not in ("كود الصنف", "اسم الصنف")}
                k = st.columns(4)
                k[0].metric("عدد الأصناف المتحرّكة", f"{len(stmt):,}")
                k[1].metric("تحميص (دخول محمص)", f"{totals.get('تحميص (دخول محمص)', 0):,.0f}")
                k[2].metric("بيع", f"{totals.get('بيع', 0):,.0f}")
                k[3].metric("فروقات جردية", f"{totals.get('فروقات جردية', 0):,.0f}")
                st.caption(f"بيان الحركات المجمّع لكل صنف — {label} (الكميات؛ الموجب وارد والسالب منصرف).")
                st.dataframe(stmt, use_container_width=True, hide_index=True)

        with st.expander("إدخال يدوي: صرف للفروع / تسوية جرد"):
            with st.form("dispatch_form", clear_on_submit=True):
                d = st.date_input("التاريخ", value=date.today(), key="disp_date")
                item = item_selector("الصنف", prefix="RB", key="disp_item")
                qty = st.number_input("الكمية المصروفة", min_value=0.0, step=1.0, key="disp_qty")
                branch = st.text_input("الفرع", key="disp_branch")
                if st.form_submit_button("💾 حفظ الصرف") and item and qty > 0:
                    add_movement(year=d.year, month=d.month, mdate=d.isoformat(), item_code=item,
                                 item_name=_name_of(item), qty=-abs(qty), movement_type=BRANCH_OUT,
                                 branch=branch, description="صرف للفرع")
                    audit("branch_out", f"{item} x{qty} -> {branch}")
                    st.success("تم تسجيل الصرف.")

            with st.form("adjust_form", clear_on_submit=True):
                d2 = st.date_input("التاريخ", value=date.today(), key="adj_date")
                item2 = item_selector("الصنف", key="adj_item")
                delta = st.number_input("قيمة التسوية (+ زيادة / - نقص)", step=1.0, key="adj_delta")
                reason = st.text_input("السبب", key="adj_reason")
                if st.form_submit_button("💾 حفظ التسوية") and item2 and delta != 0:
                    add_movement(year=d2.year, month=d2.month, mdate=d2.isoformat(), item_code=item2,
                                 item_name=_name_of(item2), qty=delta, movement_type=ADJUST,
                                 description=reason or "تسوية جرد")
                    audit("adjust", f"{item2} {delta:+g} ({reason})")
                    st.success("تم تسجيل التسوية.")

    with tab3:
        inv = movements_df()
        if inv.empty:
            st.info("لا توجد حركات.")
        else:
            types = ["الكل"] + sorted(inv["نوع الحركة"].dropna().unique().tolist())
            pick = st.selectbox("تصفية حسب نوع الحركة", types)
            view = inv if pick == "الكل" else inv[inv["نوع الحركة"] == pick]
            st.dataframe(view.tail(500), use_container_width=True, hide_index=True)


# ---- صفحة الموازنة والتوقعات ----
def page_budget():
    st.header("💰 الموازنة والتوقعات")
    settings = get_settings()
    default_growth = s_float(settings, "growth_rate")

    yrs = query_df("SELECT DISTINCT year FROM pnl_monthly ORDER BY year")
    if yrs.empty:
        st.info("لا توجد قائمة دخل شهرية بعد. ارفع ملفات الأشهر/الأعوام (Current financial year by month) "
                "ثم استورد من صفحة «📥 استيراد البيانات».")
        return
    years = [int(y) for y in yrs["year"].tolist()]

    st.caption("تُبنى الموازنة من الفعلي الشهري لكل عام/شهر ترفعه. اختر سنة الأساس وسنة الموازنة "
               "ونسبة النمو ثم ولّد؛ تُنشأ الموازنة لكل حساب ولكل شهر (أساس شهري) ويمكن عرضها سنويًا.")
    c1, c2, c3 = st.columns(3)
    base_year = c1.selectbox("سنة الأساس (الفعلي)", years, index=len(years) - 1)
    budget_year = c2.number_input("سنة الموازنة", min_value=2000, max_value=2100,
                                  value=int(base_year) + 1, step=1)
    growth_pct = c3.number_input("نسبة النمو %", value=float(default_growth * 100), step=1.0)
    growth = growth_pct / 100.0

    if st.button("🔄 توليد/تحديث الموازنة الشهرية"):
        generated = build_budget_from_monthly(base_year, growth, int(budget_year))
        execute("DELETE FROM budget WHERE year=?", (int(budget_year),))
        if not generated.empty:
            seq = [
                (int(r["السنة"]), int(r["الشهر"]), str(r["كود الحساب"]), str(r["اسم الحساب"]),
                 str(r["نوع الحساب"]), float(r["actual_prev"]), float(r["growth"]),
                 float(r["proposed"]), float(r["approved"]))
                for _, r in generated.iterrows()
            ]
            executemany(
                "INSERT INTO budget(year, month, account_code, account_name, account_type, "
                "actual_prev, growth, proposed, approved) VALUES(?,?,?,?,?,?,?,?,?)", seq,
            )
        audit("budget_generate", f"{len(generated)} rows y={budget_year} base={base_year}")
        st.success(f"تم توليد {len(generated)} بند موازنة لعام {int(budget_year)} على أساس {int(base_year)}.")
        rerun()

    budget = query_df(
        "SELECT id, year, month AS 'الشهر', account_code AS 'كود الحساب', account_name AS 'اسم الحساب', "
        "actual_prev AS 'فعلي الأساس', growth AS 'النمو', proposed AS 'موازنة مقترحة', "
        "approved AS 'موازنة معتمدة' FROM budget ORDER BY year, account_code, month"
    )
    if budget.empty:
        st.info("اضغط زر التوليد أعلاه لإنشاء الموازنة.")
    else:
        byears = sorted(int(y) for y in budget["year"].unique())
        vy = st.selectbox("سنة الموازنة المعروضة", byears, index=len(byears) - 1, key="budget_view_year")
        b = budget[budget["year"] == vy].drop(columns=["year"]).copy()

        basis = st.radio("أساس العرض", ["سنوي (ملخص)", "شهري (تفصيل وتحرير)"], horizontal=True)

        if basis.startswith("سنوي"):
            agg = b.groupby(["كود الحساب", "اسم الحساب"], as_index=False)[
                ["فعلي الأساس", "موازنة مقترحة", "موازنة معتمدة"]].sum()
            m1, m2, m3 = st.columns(3)
            m1.metric("إجمالي الفعلي (الأساس)", f"{agg['فعلي الأساس'].sum():,.0f}")
            m2.metric("إجمالي الموازنة المقترحة", f"{agg['موازنة مقترحة'].sum():,.0f}")
            m3.metric("إجمالي الموازنة المعتمدة", f"{agg['موازنة معتمدة'].sum():,.0f}")
            st.dataframe(agg, use_container_width=True, hide_index=True)
        else:
            months = sorted(int(x) for x in b["الشهر"].unique())
            mo = st.selectbox("الشهر", months,
                              format_func=lambda x: f"{x} - {MONTH_NAMES_AR[x - 1]}")
            bm = b[b["الشهر"] == mo].drop(columns=["الشهر"]).copy()
            st.subheader(f"موازنة {MONTH_NAMES_AR[mo - 1]} {int(vy)} — عدّل «النمو» و«موازنة معتمدة» ثم احفظ")
            edited = editable_table(bm, key=f"budget_editor_{vy}_{mo}")
            if edited is not None:
                if st.button("💾 حفظ تعديلات الشهر"):
                    for _, r in edited.iterrows():
                        new_growth = float(r["النمو"])
                        proposed = float(r["فعلي الأساس"]) * (1 + new_growth)
                        execute(
                            "UPDATE budget SET growth=?, proposed=?, approved=? WHERE id=?",
                            (new_growth, proposed, float(r["موازنة معتمدة"]), int(r["id"])),
                        )
                    audit("budget_edit", f"{len(edited)} rows")
                    st.success("تم حفظ التعديلات.")
                    rerun()
            else:
                st.caption("وضع التحرير المبسّط (اختر بندًا وعدّله). لجدول تحرير كامل حدّث Streamlit.")
                ids = bm["id"].tolist()
                labels = {int(r["id"]): f"{r['كود الحساب']} - {r['اسم الحساب']}" for _, r in bm.iterrows()}
                pick = st.selectbox("اختر بندًا", ids, format_func=lambda i: labels.get(int(i), str(i)))
                row = bm[bm["id"] == pick].iloc[0]
                cc1, cc2 = st.columns(2)
                new_growth = cc1.number_input("النمو", value=float(row["النمو"]), step=0.01, format="%.2f")
                new_approved = cc2.number_input("موازنة معتمدة", value=float(row["موازنة معتمدة"]), step=1.0)
                if st.button("💾 حفظ البند"):
                    proposed = float(row["فعلي الأساس"]) * (1 + new_growth)
                    execute(
                        "UPDATE budget SET growth=?, proposed=?, approved=? WHERE id=?",
                        (new_growth, proposed, new_approved, int(pick)),
                    )
                    audit("budget_edit", f"id={pick}")
                    st.success("تم حفظ البند.")
                    rerun()

    st.divider()
    st.subheader("التوقعات (متوسط آخر 6 / 12 شهرًا من الفعلي الشهري)")
    fc = build_forecast_monthly()
    if fc.empty:
        st.info("لا توجد بيانات كافية للتوقع.")
    else:
        st.dataframe(fc, use_container_width=True, hide_index=True)


# ---- صفحة الاستيراد ----
def page_import():
    st.header("📥 استيراد البيانات")
    st.caption("استورد ملفاتك الحالية إلى قاعدة البيانات. يُعاد استخدام نفس منطق قراءة ملفاتك.")

    st.subheader("1) استيراد شامل")
    st.write(
        "يقرأ تقرير حركة الصنف (الرصيد الختامي = المرجع المعتمد للأرصدة)، "
        "وتقرير المبيعات اليومية لكل فرع (عمود Branches)، وشجرة الحسابات — "
        f"من المجلد الرئيسي، بالإضافة لأي مجلدات أشهر داخل `{INPUT_DIR.name}`.\n\n"
        "ملاحظة: يُعاد الاستيراد دون تكرار، ويحافظ على أي إدخالات يدوية سجّلتها."
    )
    if st.button("🚀 استيراد كل شيء الآن"):
        with st.spinner("جارٍ الاستيراد..."):
            log = import_from_folders()
        st.success("اكتمل الاستيراد.")
        for line in log:
            st.write("•", line)

    st.divider()
    st.subheader("2) رفع ملف مفرد")
    up = st.file_uploader("ارفع ملف Excel (أرباح وخسائر) أو CSV (حركة مخزون)", type=["xlsx", "xls", "csv"])
    if up is not None:
        coly, colm = st.columns(2)
        year = coly.number_input("السنة", min_value=2000, max_value=2100, value=2025)
        month = colm.number_input("الشهر", min_value=1, max_value=12, value=12)
        if st.button("📥 استيراد الملف المرفوع"):
            tmp = DB_DIR / ("upload_" + up.name)
            with open(tmp, "wb") as f:
                f.write(up.getbuffer())
            if tmp.suffix.lower() == ".csv":
                inv = parse_inventory(tmp, int(year), int(month))
                n = insert_movements(inv)
                st.success(f"تم استيراد {n} حركة مخزون.")
            else:
                data = read_excel(tmp)
                chart, _ = load_chart_of_accounts()
                parsed = parse_profit_and_loss(data, int(year), chart)
                n = insert_financial(parsed, "pnl")
                st.success(f"تم استيراد {n} بند أرباح وخسائر.")
            try:
                tmp.unlink()
            except Exception:
                pass

    st.divider()
    st.subheader("3) استيراد شجرة الحسابات")
    if st.button("📚 تحميل شجرة الحسابات من مجلد المصادر"):
        chart, name = load_chart_of_accounts()
        if chart.empty:
            st.warning("لم يُعثر على ملف شجرة حسابات في مجلد المصادر.")
        else:
            for _, r in chart.iterrows():
                upsert_account(r["كود الحساب"], r["اسم الحساب"], r["نوع الحساب"])
            st.success(f"تم تحميل {len(chart)} حساب من {name}.")


# ---- صفحة الأصناف والحسابات ----
def page_masters():
    st.header("📚 الأصناف والحسابات")
    tab1, tab2 = st.tabs(["الأصناف", "شجرة الحسابات"])

    with tab1:
        with st.form("item_form", clear_on_submit=True):
            c1, c2 = st.columns(2)
            code = c1.text_input("كود الصنف (يبدأ بـ GB للأخضر أو RB للمحمص)")
            name = c2.text_input("اسم الصنف")
            if st.form_submit_button("➕ إضافة / تحديث صنف") and code:
                upsert_item(code, name)
                audit("item_add", code)
                st.success(f"تم حفظ الصنف {code.upper()}.")
        st.dataframe(list_items(), use_container_width=True, hide_index=True)

    with tab2:
        with st.form("acc_form", clear_on_submit=True):
            c1, c2, c3 = st.columns(3)
            code = c1.text_input("كود الحساب")
            name = c2.text_input("اسم الحساب")
            atype = c3.text_input("نوع الحساب")
            if st.form_submit_button("➕ إضافة / تحديث حساب") and code:
                upsert_account(code, name, atype)
                audit("account_add", code)
                st.success(f"تم حفظ الحساب {code}.")
        st.dataframe(list_accounts(), use_container_width=True, hide_index=True)


# ---- صفحة التنبيهات ----
def page_alerts():
    st.header("🔔 مركز التنبيهات")
    df = query_df(
        "SELECT id, created_at AS 'التاريخ', severity AS 'الأهمية', category AS 'الفئة', "
        "message AS 'الرسالة', resolved AS 'محلول' FROM alerts ORDER BY resolved, id DESC"
    )
    if df.empty:
        st.success("لا توجد تنبيهات. 🎉")
        return
    open_df = df[df["محلول"] == 0]
    st.subheader(f"تنبيهات مفتوحة ({len(open_df)})")
    for _, r in open_df.iterrows():
        c1, c2 = st.columns([5, 1])
        c1.warning(f"[{r['الفئة']}] {r['الرسالة']}  — {r['التاريخ']}")
        if c2.button("تم الحل", key=f"resolve_{r['id']}"):
            execute("UPDATE alerts SET resolved=1 WHERE id=?", (int(r["id"]),))
            audit("alert_resolved", str(r["id"]))
            rerun()
    with st.expander("عرض كل التنبيهات (المحلولة أيضًا)"):
        st.dataframe(df, use_container_width=True, hide_index=True)


# ---- صفحة الإعدادات ----
def page_settings():
    st.header("⚙️ الإعدادات")
    settings = get_settings()
    with st.form("settings_form"):
        c1, c2 = st.columns(2)
        growth = c1.number_input("نسبة النمو الافتراضية للموازنة", value=s_float(settings, "growth_rate"),
                                 step=0.01, format="%.2f")
        roast_yield = c2.number_input("نسبة إنتاج التحميص المستهدفة", value=s_float(settings, "roast_yield"),
                                      step=0.001, format="%.3f")
        tol = c1.number_input("التفاوت المسموح في التحميص", value=s_float(settings, "roast_tolerance"),
                              step=0.001, format="%.3f")
        base_year = c2.number_input("سنة الأساس", value=s_int(settings, "base_year"), step=1)
        budget_year = c1.number_input("سنة الموازنة", value=s_int(settings, "budget_year"), step=1)
        if st.form_submit_button("💾 حفظ الإعدادات"):
            set_setting("growth_rate", growth)
            set_setting("roast_yield", roast_yield)
            set_setting("roast_tolerance", tol)
            set_setting("base_year", int(base_year))
            set_setting("budget_year", int(budget_year))
            audit("settings_update", "")
            st.success("تم حفظ الإعدادات.")

    st.divider()
    st.subheader("سجل التدقيق (آخر 100 عملية)")
    st.dataframe(
        query_df("SELECT ts AS 'الوقت', usr AS 'المستخدم', action AS 'العملية', details AS 'التفاصيل' "
                 "FROM audit_log ORDER BY id DESC LIMIT 100"),
        use_container_width=True, hide_index=True,
    )


# ---- صفحة التصدير ----
def page_export():
    st.header("📤 تصدير Excel")
    st.caption("يُنشئ ملفات Excel بنفس تنسيق نظامك السابق (RTL، رؤوس ملوّنة، تجميد الصف الأول).")
    if st.button("🖨️ إنشاء ملفات الداشبورد والتنبيهات"):
        with st.spinner("جارٍ إنشاء الملفات..."):
            dash, alerts = build_exports()
        st.success("تم إنشاء الملفات في مجلد المخرجات.")
        with open(dash, "rb") as f:
            st.download_button("⬇️ تنزيل الداشبورد", f.read(), file_name=dash.name)
        with open(alerts, "rb") as f:
            st.download_button("⬇️ تنزيل التنبيهات", f.read(), file_name=alerts.name)


# ---- صفحة مبيعات التطبيقات والموقع (إدخال يدوي شهري) ----
MONTH_NAMES_AR = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
                  "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]


def page_channels():
    st.header("🛵 مبيعات التطبيقات والموقع الإلكتروني")
    st.caption("أدخل مبيعات كل قناة (جاهز، نوفمبر، هنجرستيشن، الموقع الإلكتروني) شهريًا. "
               "إعادة الإدخال لنفس (السنة/الشهر/القناة) تُحدِّث القيمة.")

    with st.form("channel_form", clear_on_submit=True):
        c1, c2, c3 = st.columns(3)
        year = c1.number_input("السنة", min_value=2000, max_value=2100, value=2025, step=1)
        month = c2.selectbox("الشهر", list(range(1, 13)),
                             format_func=lambda m: f"{m} - {MONTH_NAMES_AR[m - 1]}")
        channel = c3.selectbox("القناة", DELIVERY_CHANNELS)
        custom = st.text_input("اسم قناة مخصص (اختياري، يُستخدم بدل الاختيار أعلاه)")
        amount = st.number_input("قيمة المبيعات", min_value=0.0, step=100.0)
        note = st.text_input("ملاحظة (اختياري)")
        if st.form_submit_button("💾 حفظ / تحديث"):
            ch = custom.strip() or channel
            if amount > 0:
                upsert_channel_sale(int(year), int(month), ch, amount, note)
                audit("channel_sale", f"{year}-{month} {ch}={amount}")
                st.success(f"تم حفظ مبيعات {ch} لشهر {MONTH_NAMES_AR[int(month) - 1]} {int(year)}.")
            else:
                st.error("أدخل قيمة أكبر من صفر.")

    df = channel_sales_df()
    if df.empty:
        st.info("لا توجد مبيعات قنوات مُدخلة بعد.")
        return

    st.divider()
    years = sorted(df["السنة"].unique())
    ysel = st.selectbox("اعرض سنة", years, index=len(years) - 1, key="ch_year")
    year_df = df[df["السنة"] == ysel]
    st.subheader(f"مبيعات القنوات لعام {int(ysel)}")
    total = year_df["المبلغ"].sum()
    st.metric("إجمالي مبيعات القنوات", f"{total:,.0f}")
    pivot = year_df.pivot_table(index="الشهر", columns="القناة", values="المبلغ",
                                aggfunc="sum", fill_value=0)
    st.dataframe(pivot, use_container_width=True)
    by_channel = year_df.groupby("القناة")["المبلغ"].sum().sort_values(ascending=False)
    st.bar_chart(by_channel)
    with st.expander("كل السجلات"):
        st.dataframe(df, use_container_width=True, hide_index=True)


# ---- صفحة التحليل المالي والنسب المحاسبية ----
def _fmt_compare_table(metric_dict_by_year, years):
    rows = []
    metrics = list(next(iter(metric_dict_by_year.values())).keys())
    for metric in metrics:
        row = {"البند": metric}
        for y in years:
            row[str(int(y))] = round(metric_dict_by_year[y][metric], 2)
        if len(years) == 2:
            a, b = metric_dict_by_year[years[0]][metric], metric_dict_by_year[years[1]][metric]
            row["التغير %"] = round((b - a) / a * 100, 1) if a else 0.0
        rows.append(row)
    return pd.DataFrame(rows)


def page_finance():
    st.header("📊 التحليل المالي والنسب المحاسبية")
    years = list_years()
    if not years:
        st.info("لا توجد بيانات مالية بعد. استورد قائمة الدخل والميزانية من صفحة «استيراد البيانات».")
        return

    tab1, tab2, tab3 = st.tabs(["مقارنة 12 شهر", "الربحية (سنوي)", "النسب المالية"])

    # ---------- Tab 1: 12-month comparison ----------
    with tab1:
        st.caption("مقارنة المبيعات والمصاريف شهريًا بين سنتين (المبيعات = قائمة الدخل + مبيعات القنوات اليدوية).")
        c1, c2 = st.columns(2)
        ya = c1.selectbox("السنة الأولى", years, index=0, key="fa_ya")
        yb = c2.selectbox("السنة الثانية", years, index=len(years) - 1, key="fa_yb")

        # المبيعات: قائمة الدخل + القنوات اليدوية
        def revenue_of(y):
            base = monthly_pnl_series(y, "revenue")
            ch = pd.Series(0.0, index=range(1, 13))
            cq = query_df("SELECT month, COALESCE(SUM(amount),0) v FROM channel_sales WHERE year=? GROUP BY month", (int(y),))
            for _, r in cq.iterrows():
                m = int(r["month"])
                if 1 <= m <= 12:
                    ch[m] += float(r["v"])
            return base.add(ch, fill_value=0)

        def expenses_of(y):
            return monthly_pnl_series(y, "cogs").add(monthly_pnl_series(y, "expense"), fill_value=0)

        rev_a, rev_b = revenue_of(ya), revenue_of(yb)
        exp_a, exp_b = expenses_of(ya), expenses_of(yb)
        table = pd.DataFrame({
            "الشهر": MONTH_NAMES_AR,
            f"مبيعات {int(ya)}": rev_a.values, f"مبيعات {int(yb)}": rev_b.values,
            f"مصاريف {int(ya)}": exp_a.values, f"مصاريف {int(yb)}": exp_b.values,
        })
        table[f"صافي {int(ya)}"] = table[f"مبيعات {int(ya)}"] - table[f"مصاريف {int(ya)}"]
        table[f"صافي {int(yb)}"] = table[f"مبيعات {int(yb)}"] - table[f"مصاريف {int(yb)}"]
        totals = {"الشهر": "الإجمالي"}
        for c in table.columns[1:]:
            totals[c] = table[c].sum()
        show = pd.concat([table, pd.DataFrame([totals])], ignore_index=True)
        for c in show.columns[1:]:
            show[c] = show[c].round(0)
        st.dataframe(show, use_container_width=True, hide_index=True)

        st.subheader("المبيعات شهريًا")
        st.line_chart(pd.DataFrame({int(ya): rev_a.values, int(yb): rev_b.values},
                                   index=MONTH_NAMES_AR))
        st.subheader("المصاريف شهريًا")
        st.line_chart(pd.DataFrame({int(ya): exp_a.values, int(yb): exp_b.values},
                                   index=MONTH_NAMES_AR))
        d_rev = rev_b.sum() - rev_a.sum()
        d_exp = exp_b.sum() - exp_a.sum()
        k1, k2, k3 = st.columns(3)
        k1.metric(f"إجمالي مبيعات {int(yb)}", f"{rev_b.sum():,.0f}",
                  f"{(d_rev / rev_a.sum() * 100) if rev_a.sum() else 0:,.1f}%")
        k2.metric(f"إجمالي مصاريف {int(yb)}", f"{exp_b.sum():,.0f}",
                  f"{(d_exp / exp_a.sum() * 100) if exp_a.sum() else 0:,.1f}%", delta_color="inverse")
        k3.metric(f"صافي {int(yb)}", f"{(rev_b.sum() - exp_b.sum()):,.0f}")

    # ---------- Tab 2: annual profitability ----------
    with tab2:
        st.caption("ملخص الربحية السنوي من قائمة الدخل.")
        summ = {y: pnl_year_summary(y) for y in years}
        if all(all(v == 0 for v in s.values()) for s in summ.values()):
            st.warning("لا توجد بيانات قائمة دخل شهرية. تأكد من استيراد ملف «Current financial year by month».")
        else:
            comp = _fmt_compare_table(summ, years)
            st.dataframe(comp, use_container_width=True, hide_index=True)
            latest = years[-1]
            s = summ[latest]
            c1, c2, c3, c4 = st.columns(4)
            c1.metric(f"إيرادات {int(latest)}", f"{s['الإيرادات']:,.0f}")
            c2.metric("مجمل الربح", f"{s['مجمل الربح']:,.0f}")
            c3.metric("صافي الربح", f"{s['صافي الربح']:,.0f}")
            c4.metric("هامش صافي الربح", f"{s['هامش صافي الربح %']:,.1f}%")

    # ---------- Tab 3: financial ratios ----------
    with tab3:
        st.caption("النسب المالية (السيولة، دوران المخزون، الهوامش، العائد). "
                   "أرقام الميزانية تُقرأ آليًا من ملف الميزانية، ويمكنك تعديلها يدويًا أدناه.")
        ysel = st.selectbox("السنة", years, index=len(years) - 1, key="ratio_year")
        ratios, figures = financial_ratios(ysel)

        rt = pd.DataFrame({"النسبة": list(ratios.keys()),
                           "القيمة": [round(v, 2) for v in ratios.values()]})
        fg = pd.DataFrame({"البند": list(figures.keys()),
                           "القيمة": [round(v, 0) for v in figures.values()]})
        cc1, cc2 = st.columns(2)
        cc1.subheader("النسب")
        cc1.dataframe(rt, use_container_width=True, hide_index=True)
        cc2.subheader("أرقام الميزانية المستخدمة")
        cc2.dataframe(fg, use_container_width=True, hide_index=True)

        with st.expander("تعديل أرقام الميزانية يدويًا (اختياري)"):
            with st.form("fig_form"):
                f1, f2, f3 = st.columns(3)
                ca = f1.number_input("الأصول المتداولة", value=float(figures["الأصول المتداولة"]), step=1000.0)
                inv = f2.number_input("المخزون", value=float(figures["المخزون"]), step=1000.0)
                cl = f3.number_input("الخصوم المتداولة", value=float(figures["الخصوم المتداولة"]), step=1000.0)
                ta = f1.number_input("إجمالي الأصول", value=float(figures["إجمالي الأصول"]), step=1000.0)
                te = f2.number_input("حقوق الملكية", value=float(figures["حقوق الملكية"]), step=1000.0)
                if st.form_submit_button("💾 حفظ الأرقام"):
                    upsert_fin_figures(pd.DataFrame([{
                        "year": int(ysel), "current_assets": ca, "inventory": inv,
                        "current_liabilities": cl, "total_assets": ta, "total_equity": te,
                        "source_file": "إدخال يدوي",
                    }]))
                    audit("fin_figures_manual", str(ysel))
                    st.success("تم الحفظ. أعد اختيار السنة لتحديث النسب.")


# ==========================================================================
# التشغيل
# ==========================================================================
def main():
    setup_page()
    init_db()
    page = sidebar()
    if page.startswith("🏠"):
        page_dashboard()
    elif page.startswith("🧾"):
        page_sales()
    elif page.startswith("🛵"):
        page_channels()
    elif page.startswith("🔥"):
        page_roasting()
    elif page.startswith("📦"):
        page_inventory()
    elif page.startswith("💰"):
        page_budget()
    elif page.startswith("📊"):
        page_finance()
    elif page.startswith("📥"):
        page_import()
    elif page.startswith("📚"):
        page_masters()
    elif page.startswith("🔔"):
        page_alerts()
    elif page.startswith("⚙️"):
        page_settings()
    elif page.startswith("📤"):
        page_export()


if __name__ == "__main__":
    main()