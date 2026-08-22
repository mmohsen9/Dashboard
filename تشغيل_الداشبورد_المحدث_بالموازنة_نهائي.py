from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

import pandas as pd

BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "01_ملفات_الإدخال_الشهرية"
STATIC_DIR = BASE_DIR / "02_مصادر_ثابتة"
OUTPUT_DIR = BASE_DIR / "04_المخرجات"

DASHBOARD_PATH = OUTPUT_DIR / "الداشبورد_المالي_والتشغيلي.xlsx"
ALERTS_PATH = OUTPUT_DIR / "تنبيهات_الشهر.xlsx"

GROWTH_RATE = 0.10
ROAST_YIELD = 0.833
ROAST_TOLERANCE = 0.012


def normalize(value):
    text = str(value).replace("\n", " ").replace("\r", " ").strip().lower()
    return re.sub(r"\s+", " ", text)


def clean_code(value):
    if pd.isna(value):
        return ""
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") else text


def number(value):
    if pd.isna(value):
        return 0.0
    text = str(value).strip().replace(",", "")
    text = text.replace("(", "-").replace(")", "")
    result = pd.to_numeric(text, errors="coerce")
    return float(result) if pd.notna(result) else 0.0


def find_column(columns, words):
    for column in columns:
        if any(word in normalize(column) for word in words):
            return column
    return None


def read_excel(path):
    try:
        return pd.read_excel(path, sheet_name=0)
    except Exception:
        return pd.DataFrame()


def read_csv_rows(path):
    for encoding in ["utf-8-sig", "utf-8", "cp1256", "latin1"]:
        try:
            with open(path, "r", encoding=encoding, errors="replace", newline="") as file:
                return list(csv.reader(file, delimiter=";"))
        except Exception:
            continue
    return []


def load_chart_of_accounts():
    files = list(STATIC_DIR.glob("*.xlsx")) + list(BASE_DIR.glob("*.xlsx"))

    for path in files:
        name = path.name.lower()
        if "chart" not in name and "account" not in name:
            continue

        data = read_excel(path)
        if data.empty:
            continue

        code_column = find_column(data.columns, ["account code", "رقم الحساب", "كود الحساب"])
        name_column = find_column(data.columns, ["account اسم", "account name", "account", "اسم الحساب"])
        type_column = find_column(data.columns, ["account type", "نوع الحساب"])

        if code_column is None or name_column is None:
            continue

        chart = pd.DataFrame({
            "كود الحساب": data[code_column].map(clean_code),
            "اسم الحساب": data[name_column].astype(str).str.strip(),
            "نوع الحساب": data[type_column].astype(str).str.strip() if type_column else ""
        })

        chart = chart[chart["كود الحساب"] != ""].drop_duplicates("كود الحساب")
        print("تم تحميل شجرة الحسابات:", path.name)
        print("عدد الحسابات:", len(chart))
        return chart

    print("تحذير: لم يتم العثور على شجرة الحسابات")
    return pd.DataFrame(columns=["كود الحساب", "اسم الحساب", "نوع الحساب"])


def attach_accounts(data, chart, code_column=None, name_column=None):
    output = data.copy()

    chart_by_code = {}
    chart_by_name = {}

    for _, row in chart.iterrows():
        account_code = clean_code(row["كود الحساب"])
        account_name = str(row["اسم الحساب"]).strip()
        record = {
            "كود الحساب": account_code,
            "اسم الحساب": account_name,
            "نوع الحساب": str(row["نوع الحساب"]).strip()
        }
        chart_by_code[account_code] = record
        chart_by_name[normalize(account_name)] = record

    if code_column and code_column in output.columns:
        output["كود الحساب"] = output[code_column].map(clean_code)
    else:
        output["كود الحساب"] = ""

    if name_column and name_column in output.columns:
        source_names = output[name_column].astype(str).str.strip()
    else:
        source_names = pd.Series([""] * len(output), index=output.index)

    output["اسم الحساب"] = ""
    output["نوع الحساب"] = ""

    for index in output.index:
        account_code = clean_code(output.at[index, "كود الحساب"])
        source_name = str(source_names.loc[index]).strip()
        matched = chart_by_code.get(account_code)

        if matched is None and source_name:
            matched = chart_by_name.get(normalize(source_name))

        if matched is None and source_name:
            source_key = normalize(source_name)
            for chart_name, chart_record in chart_by_name.items():
                if source_key in chart_name or chart_name in source_key:
                    matched = chart_record
                    break

        if matched is not None:
            output.at[index, "كود الحساب"] = matched["كود الحساب"]
            output.at[index, "اسم الحساب"] = matched["اسم الحساب"]
            output.at[index, "نوع الحساب"] = matched["نوع الحساب"]
        else:
            output.at[index, "اسم الحساب"] = source_name

    return output


def extract_year_month(folder):
    folder_name = str(folder)
    match = re.search(r"(20\d{2})[-_](0?[1-9]|1[0-2])", folder_name)
    return (int(match.group(1)), int(match.group(2))) if match else (2025, 12)


def parse_profit_and_loss(data, year, chart):
    if data.empty:
        return pd.DataFrame()

    account_column = find_column(data.columns, ["account", "account name", "اسم الحساب"])
    if account_column is None:
        account_column = data.columns[0]

    rows = []

    for _, row in data.iterrows():
        account_text = str(row.get(account_column, "")).strip()
        if not account_text:
            continue

        if normalize(account_text).startswith(("total", "gross profit", "net profit", "trading income", "cost of sales", "operating expenses", "other income", "إجمالي", "صافي")):
            continue

        match = re.match(r"^\s*(\d+)\s*[-–]\s*(.*)$", account_text)
        source_code = match.group(1).strip() if match else ""
        source_name = match.group(2).strip() if match else account_text

        amount = sum(
            number(row.get(column, 0))
            for column in data.columns
            if column != account_column
        )

        if amount:
            rows.append({
                "السنة": year,
                "الشهر": 12,
                "كود_مصدر": source_code,
                "اسم_مصدر": source_name,
                "المبلغ": amount
            })

    result = pd.DataFrame(rows)
    if result.empty:
        return result

    result = attach_accounts(result, chart, "كود_مصدر", "اسم_مصدر")

    return result[["السنة", "الشهر", "كود الحساب", "اسم الحساب", "نوع الحساب", "المبلغ"]]


def parse_trial_or_balance(data, year, month, chart):
    if data.empty:
        return pd.DataFrame()

    code_column = find_column(data.columns, ["account code", "رقم الحساب", "كود الحساب"])
    name_column = find_column(data.columns, ["account name", "account", "اسم الحساب"])
    result = attach_accounts(data, chart, code_column, name_column)
    result["السنة"] = year
    result["الشهر"] = month
    return result


def parse_inventory(path, year, month):
    rows = read_csv_rows(path)
    if not rows:
        return pd.DataFrame()

    header_index = None
    for index, row in enumerate(rows):
        text = normalize(" ".join(row))
        if "date" in text and "qoh movement" in text:
            header_index = index
            break

    if header_index is None:
        return pd.DataFrame()

    headers = rows[header_index]
    records = []
    current_code = ""
    current_name = ""

    for row in rows[header_index + 1:]:
        if len(row) < len(headers):
            row += [""] * (len(headers) - len(row))

        values = dict(zip(headers, row[:len(headers)]))
        first_cell = str(row[0]).strip() if row else ""

        item_match = re.match(r"^\s*((?:GB|RB)[A-Z0-9]*)\s*[-–]\s*(.*)", first_cell, re.IGNORECASE)
        if item_match:
            current_code = item_match.group(1).upper()
            current_name = item_match.group(2).strip()
            continue

        date_value = values.get("Date", "")
        if not str(date_value).strip():
            continue

        text = normalize(
            f"{values.get('Source', '')} "
            f"{values.get('Reference', '')} "
            f"{values.get('Contact', '')} "
            f"{values.get('Description', '')}"
        )

        if "امر تحميص مستودع البن الاخضر" in text:
            movement_type = "خروج أخضر للتحميص"
        elif "تحميص وارد لمستودع البن المحمص" in text:
            movement_type = "دخول محمص من التحميص"
        elif "سند صرف قهوه محموصه فرع" in text:
            movement_type = "صرف محمص للفروع"
        elif "عينات القهوة المحمصة" in text:
            movement_type = "عينات وهدايا"
        elif "receivable invoice" in text:
            movement_type = "بيع"
        elif any(word in text for word in ["adjustment", "جرد", "revaluation", "تعديل"]):
            movement_type = "تسوية أو إعادة تقييم"
        else:
            movement_type = "أخرى"

        coffee_type = "بن أخضر" if current_code.startswith("GB") else "بن محمص" if current_code.startswith("RB") else "غير مصنف"
        family_code = re.sub(r"^(GB|RB)", "", current_code, flags=re.IGNORECASE)

        records.append({
            "السنة": year,
            "الشهر": month,
            "التاريخ": date_value,
            "كود الصنف": current_code,
            "اسم الصنف": current_name,
            "نوع البن": coffee_type,
            "رقم العائلة": family_code,
            "المصدر": values.get("Source", ""),
            "المرجع": values.get("Reference", ""),
            "العميل أو المورد": values.get("Contact", ""),
            "الوصف": values.get("Description", ""),
            "الإجمالي": number(values.get("Total", 0)),
            "قيمة الحركة": number(values.get("Value Movement", 0)),
            "حركة الكمية": number(values.get("QoH Movement", 0)),
            "حساب الإيراد": values.get("Revenue Account", ""),
            "حساب المخزون": values.get("Inventory Account", ""),
            "نوع الحركة": movement_type
        })

    return pd.DataFrame(records)


def build_budget(pnl):
    if pnl.empty:
        return pd.DataFrame()

    result = pnl[pnl["السنة"] == 2025].groupby(
        ["الشهر", "كود الحساب", "اسم الحساب", "نوع الحساب"],
        as_index=False
    )["المبلغ"].sum()

    result["السنة"] = 2026
    result["فعلي 2025"] = result["المبلغ"]
    result["نمو مقترح"] = GROWTH_RATE
    result["موازنة مقترحة"] = result["فعلي 2025"] * (1 + GROWTH_RATE)
    result["موازنة معتمدة"] = result["موازنة مقترحة"]

    return result[["السنة", "الشهر", "كود الحساب", "اسم الحساب", "نوع الحساب", "فعلي 2025", "نمو مقترح", "موازنة مقترحة", "موازنة معتمدة"]]


def build_forecast(pnl):
    if pnl.empty:
        return pd.DataFrame()

    rows = []

    for (account_code, account_name, account_type), group in pnl.groupby(["كود الحساب", "اسم الحساب", "نوع الحساب"]):
        values = group.sort_values(["السنة", "الشهر"])["المبلغ"]
        average_6 = values.tail(6).mean()
        average_12 = values.tail(12).mean()

        rows.append({
            "كود الحساب": account_code,
            "اسم الحساب": account_name,
            "نوع الحساب": account_type,
            "متوسط آخر 6 أشهر": average_6,
            "متوسط آخر 12 شهرًا": average_12,
            "توقع الشهر التالي": average_6,
            "منهج التوقع": "متوسط آخر 6 أشهر مع مقارنة متوسط 12 شهرًا"
        })

    return pd.DataFrame(rows)


def build_roasting(inventory):
    if inventory.empty:
        return pd.DataFrame()

    green = inventory[inventory["نوع الحركة"] == "خروج أخضر للتحميص"]]
    roasted = inventory[inventory["نوع الحركة"] == "دخول محمص من التحميص"]]

    green_summary = green.groupby(["السنة", "الشهر", "رقم العائلة"], as_index=False)["حركة الكمية"].sum()
    green_summary = green_summary.rename(columns={"حركة الكمية": "كمية الأخضر المحولة"})
    green_summary["كمية الأخضر المحولة"] = green_summary["كمية الأخضر المحولة"].abs()

    roasted_summary = roasted.groupby(["السنة", "الشهر", "رقم العائلة"], as_index=False)["حركة الكمية"].sum()
    roasted_summary = roasted_summary.rename(columns={"حركة الكمية": "كمية المحمص الناتجة"})

    result = green_summary.merge(roasted_summary, on=["السنة", "الشهر", "رقم العائلة"], how="outer").fillna(0)

    if result.empty:
        return result

    names = inventory[["رقم العائلة", "اسم الصنف"]].copy()
    names["اسم الصنف"] = names["اسم الصنف"].astype(str).str.strip()
    names = names[names["اسم الصنف"] != ""].drop_duplicates("رقم العائلة")
    names = names.rename(columns={"اسم الصنف": "الصنف"})
    result = result.merge(names, on="رقم العائلة", how="left")

    result["الناتج المتوقع"] = result["كمية الأخضر المحولة"] * ROAST_YIELD
    result["فرق الكمية"] = result["كمية المحمص الناتجة"] - result["الناتج المتوقع"]
    result["فرق النسبة"] = result.apply(lambda row: row["فرق الكمية"] / row["الناتج المتوقع"] if row["الناتج المتوقع"] else 0, axis=1)
    result["الحالة"] = result["فرق النسبة"].map(lambda value: "مقبول" if abs(value) <= ROAST_TOLERANCE else "يحتاج مراجعة")

    return result[["السنة", "الشهر", "رقم العائلة", "الصنف", "كمية الأخضر المحولة", "كمية المحمص الناتجة", "الناتج المتوقع", "فرق الكمية", "فرق النسبة", "الحالة"]]


def build_inventory_summary(inventory):
    if inventory.empty:
        return pd.DataFrame()

    return inventory.groupby(["السنة", "الشهر", "كود الصنف", "اسم الصنف", "نوع البن", "رقم العائلة", "نوع الحركة"], as_index=False).agg({"حركة الكمية": "sum", "قيمة الحركة": "sum", "الإجمالي": "sum"})


def write_sheet(writer, name, data):
    if data is None or data.empty:
        data = pd.DataFrame({"ملاحظة": ["لا توجد بيانات متاحة لهذه الصفحة"]})

    data.to_excel(writer, sheet_name=name[:31], index=False)


def style_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.rightToLeft = True
        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill

    workbook.save(path)


def main():
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    STATIC_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    chart = load_chart_of_accounts()
    folders = sorted([folder for folder in INPUT_DIR.iterdir() if folder.is_dir()])

    if not folders:
        raise FileNotFoundError(f"لا توجد مجلدات أشهر داخل: {INPUT_DIR}")

    pnl_parts = []
    trial_parts = []
    balance_parts = []
    inventory_parts = []

    for folder in folders:
        year, month = extract_year_month(folder)
        print("معالجة الشهر:", folder.name)

        for path in folder.iterdir():
            extension = path.suffix.lower()

            if extension not in [".xlsx", ".xls", ".csv"]:
                continue

            if extension == ".csv":
                inventory = parse_inventory(path, year, month)
                if not inventory.empty:
                    inventory_parts.append(inventory)
                print(path.name, "=> حركة المخزون:", len(inventory))
                continue

            data = read_excel(path)
            if data.empty:
                continue

            file_name = path.name.lower()
            signature = " ".join(normalize(column) for column in data.columns)

            if "profit and loss" in file_name or "profit and loss" in signature:
                parsed = parse_profit_and_loss(data, year, chart)
                if not parsed.empty:
                    pnl_parts.append(parsed)
                print(path.name, "=> P&L:", len(parsed))

            elif "trial balance" in file_name or "trial balance" in signature:
                parsed = parse_trial_or_balance(data, year, month, chart)
                if not parsed.empty:
                    trial_parts.append(parsed)
                print(path.name, "=> Trial Balance:", len(parsed))

            elif "balance sheet" in file_name or "balance sheet" in signature:
                parsed = parse_trial_or_balance(data, year, month, chart)
                if not parsed.empty:
                    balance_parts.append(parsed)
                print(path.name, "=> Balance Sheet:", len(parsed))

    pnl = pd.concat(pnl_parts, ignore_index=True) if pnl_parts else pd.DataFrame()
    trial_balance = pd.concat(trial_parts, ignore_index=True) if trial_parts else pd.DataFrame()
    balance_sheet = pd.concat(balance_parts, ignore_index=True) if balance_parts else pd.DataFrame()
    inventory_raw = pd.concat(inventory_parts, ignore_index=True) if inventory_parts else pd.DataFrame()

    budget = build_budget(pnl)
    forecast = build_forecast(pnl)
    inventory_summary = build_inventory_summary(inventory_raw)
    roasting = build_roasting(inventory_raw)

    with pd.ExcelWriter(DASHBOARD_PATH, engine="openpyxl") as writer:
        write_sheet(writer, "شجرة الحسابات", chart)
        write_sheet(writer, "قائمة الأرباح والخسائر", pnl)
        write_sheet(writer, "ميزان المراجعة", trial_balance)
        write_sheet(writer, "الميزانية العمومية", balance_sheet)
        write_sheet(writer, "2026 الموازنة", budget)
        write_sheet(writer, "التوقعات", forecast)
        write_sheet(writer, "حركة المخزون", inventory_raw)
        write_sheet(writer, "تحليل المخزون", inventory_summary)
        write_sheet(writer, "تحليل التحميص", roasting)

    style_workbook(DASHBOARD_PATH)

    alerts = roasting[roasting["الحالة"] == "يحتاج مراجعة"] if not roasting.empty else pd.DataFrame()

    with pd.ExcelWriter(ALERTS_PATH, engine="openpyxl") as writer:
        write_sheet(writer, "فروقات التحميص", alerts)
        write_sheet(writer, "ملخص المخزون", inventory_summary)

    style_workbook(ALERTS_PATH)

    print("=" * 60)
    print("تم إنشاء الملفات بنجاح")
    print("عدد حسابات الشجرة:", len(chart))
    print("صفوف الأرباح والخسائر:", len(pnl))
    print("صفوف ميزان المراجعة:", len(trial_balance))
    print("صفوف حركة المخزون:", len(inventory_raw))
    print("صفوف تحليل التحميص:", len(roasting))
    print("ملف الداشبورد:", DASHBOARD_PATH)
    print("ملف التنبيهات:", ALERTS_PATH)
    print("=" * 60)


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print("حدث خطأ أثناء التشغيل:")
        print(error)
        sys.exit(1)